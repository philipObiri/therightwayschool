"""Tests for Bulk Upload endpoints (/api/bulk-upload/preview + /create-accounts)."""
import os
import io
import pytest
import requests
import openpyxl
from docx import Document

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN = {"email": "dwain@tcis.demo", "password": "admin123"}
TEACHER = {"email": "edmund@tcis.demo", "password": "teacher123"}


# ─── helpers ─────────────────────────────────────────────────────────
def login(creds):
    r = requests.post(f"{API}/auth/login", json=creds, timeout=20)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    return r.json().get("token") or r.json().get("access_token")


def build_xlsx(rows):
    """rows = list of dicts; first row keys = headers"""
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = list(rows[0].keys())
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def build_docx(rows):
    doc = Document()
    headers = list(rows[0].keys())
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for ri, r in enumerate(rows, 1):
        for i, h in enumerate(headers):
            t.rows[ri].cells[i].text = str(r.get(h, ""))
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


SAMPLE_ROWS = [
    {
        "Student Name": "TEST_Bulk Alpha",
        "Year Group": "Year 7",
        "Gender": "Male",
        "Boarding Status": "Day",
        "Parent Name": "TEST_Bulk Parent Alpha",
        "Parent Phone": "0555000001",
        "Parent Email": "testbulk.alphaparent@example.com",
        "Address": "Accra",
    },
    {
        "Student Name": "TEST_Bulk Beta",
        "Year Group": "Year 8",
        "Gender": "Female",
        "Boarding Status": "Boarding",
        "Parent Name": "TEST_Bulk Parent Beta",
        "Parent Phone": "0555000002",
        "Parent Email": "",  # auto-generate
        "Address": "Tema",
    },
]


# ─── fixtures ─────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def admin_token():
    return login(ADMIN)


@pytest.fixture(scope="module")
def teacher_token():
    try:
        return login(TEACHER)
    except Exception:
        pytest.skip("Teacher login unavailable")


@pytest.fixture(scope="module", autouse=True)
def cleanup():
    yield
    # Clean up test records
    import asyncio
    from database import db

    async def purge():
        await db.users.delete_many({"name": {"$regex": "^TEST_Bulk"}})
        await db.users.delete_many({"email": {"$regex": "testbulk\\."}})
        await db.users.delete_many({"email": {"$regex": "^parent\\.test_bulk"}})
    asyncio.run(purge())


def auth_hdr(token):
    return {"Authorization": f"Bearer {token}"}


# ─── Preview tests ─────────────────────────────────────────────────────
class TestPreview:
    def test_preview_requires_auth(self):
        files = {"file": ("t.xlsx", build_xlsx(SAMPLE_ROWS),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/bulk-upload/preview", files=files, timeout=20)
        assert r.status_code in (401, 403)

    def test_preview_forbids_teacher(self, teacher_token):
        files = {"file": ("t.xlsx", build_xlsx(SAMPLE_ROWS),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/bulk-upload/preview", files=files,
                          headers=auth_hdr(teacher_token), timeout=20)
        assert r.status_code == 403

    def test_preview_xlsx(self, admin_token):
        files = {"file": ("t.xlsx", build_xlsx(SAMPLE_ROWS),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/bulk-upload/preview", files=files,
                          headers=auth_hdr(admin_token), timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["total_rows"] == 2
        cm = d["column_mapping"]
        # auto-detected mapping
        assert "name" in cm and "year_group" in cm and "gender" in cm
        assert "parent_email" in cm and "parent_phone" in cm and "boarding_status" in cm
        assert d["all_data"][0]["student_name"] == "TEST_Bulk Alpha"
        assert d["all_data"][0]["gender"] == "Male"
        assert d["all_data"][1]["boarding_status"] == "boarding"

    def test_preview_docx(self, admin_token):
        files = {"file": ("t.docx", build_docx(SAMPLE_ROWS),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = requests.post(f"{API}/bulk-upload/preview", files=files,
                          headers=auth_hdr(admin_token), timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["total_rows"] == 2
        assert d["all_data"][0]["student_name"] == "TEST_Bulk Alpha"

    def test_preview_rejects_txt(self, admin_token):
        files = {"file": ("t.txt", b"hello world", "text/plain")}
        r = requests.post(f"{API}/bulk-upload/preview", files=files,
                          headers=auth_hdr(admin_token), timeout=20)
        assert r.status_code == 400


# ─── Create-accounts tests ─────────────────────────────────────────────
class TestCreateAccounts:
    def test_create_and_parent_login(self, admin_token):
        # 1. preview
        files = {"file": ("t.xlsx", build_xlsx(SAMPLE_ROWS),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        pv = requests.post(f"{API}/bulk-upload/preview", files=files,
                           headers=auth_hdr(admin_token), timeout=30).json()

        # 2. create
        r = requests.post(f"{API}/bulk-upload/create-accounts",
                          json={"records": pv["all_data"], "create_students": True, "create_parents": True},
                          headers=auth_hdr(admin_token), timeout=60)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["students_created"] == 2
        assert d["parents_created"] == 2
        assert d["default_passwords"]["parents"] == "parent123"
        assert d["default_passwords"]["students"] == "student123"

        # 3. parent (explicit email) can login
        lr = requests.post(f"{API}/auth/login",
                           json={"email": "testbulk.alphaparent@example.com", "password": "parent123"},
                           timeout=20)
        assert lr.status_code == 200, lr.text
        assert lr.json()["user"]["role"] == "Parent"

        # 4. auto-generated parent email for beta exists & can login
        auto_email = "parent.test_bulk.beta@tcis.edu.gh"
        lr2 = requests.post(f"{API}/auth/login",
                            json={"email": auto_email, "password": "parent123"},
                            timeout=20)
        assert lr2.status_code == 200, f"auto-gen parent login failed: {lr2.status_code} {lr2.text}"

    def test_duplicate_skip_and_parent_update(self, admin_token):
        # re-upload same data: students already exist -> skipped, parents exist -> updated
        files = {"file": ("t.xlsx", build_xlsx(SAMPLE_ROWS),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        pv = requests.post(f"{API}/bulk-upload/preview", files=files,
                           headers=auth_hdr(admin_token), timeout=30).json()
        r = requests.post(f"{API}/bulk-upload/create-accounts",
                          json={"records": pv["all_data"]},
                          headers=auth_hdr(admin_token), timeout=60)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["students_skipped"] == 2
        assert d["students_created"] == 0
        assert d["parents_updated"] == 2
        assert d["parents_created"] == 0

    def test_create_students_false(self, admin_token):
        rows = [{
            "Student Name": "TEST_Bulk Gamma",
            "Year Group": "Year 9",
            "Parent Name": "TEST_Bulk Parent Gamma",
            "Parent Email": "testbulk.gammaparent@example.com",
            "Parent Phone": "0555000003",
        }]
        files = {"file": ("t.xlsx", build_xlsx(rows),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        pv = requests.post(f"{API}/bulk-upload/preview", files=files,
                           headers=auth_hdr(admin_token), timeout=30).json()
        r = requests.post(f"{API}/bulk-upload/create-accounts",
                          json={"records": pv["all_data"], "create_students": False, "create_parents": True},
                          headers=auth_hdr(admin_token), timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["students_created"] == 0
        assert d["parents_created"] == 1

    def test_empty_records_400(self, admin_token):
        r = requests.post(f"{API}/bulk-upload/create-accounts",
                          json={"records": []},
                          headers=auth_hdr(admin_token), timeout=20)
        assert r.status_code == 400


# ─── regression: existing credentials still work ─────────────────────
class TestRegression:
    def test_admin_login(self):
        t = login(ADMIN)
        assert t

    def test_existing_parent_login(self):
        r = requests.post(f"{API}/auth/login",
                          json={"email": "sheilarthur26@gmail.com", "password": "parent123"},
                          timeout=20)
        assert r.status_code == 200, r.text
        assert r.json()["user"]["role"] == "Parent"
