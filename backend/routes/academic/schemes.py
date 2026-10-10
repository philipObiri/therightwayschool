from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from database import db, get_active_term
from auth import get_current_user
from routes.approval_chain import (
    APPROVAL_CHAIN, build_approval_chain, get_current_step_index,
    is_chain_complete, can_user_approve_step, get_current_step_label,
    can_user_see_lesson,
)
from datetime import datetime, timezone, timedelta
import jwt as pyjwt
import uuid
import os
import re
import base64

JWT_SECRET = os.environ.get("JWT_SECRET", "default_secret")
JWT_ALGORITHM = "HS256"

async def _get_user_from_token_or_header(request: Request, token: str = None):
    """Get user from query param token or Authorization header."""
    if token:
        try:
            payload = pyjwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
            user_id = payload.get("user_id")
            if user_id:
                user = await db.users.find_one({"user_id": user_id}, {"_id": 0, "password_hash": 0})
                if user:
                    return user
        except Exception:
            pass
    return await get_current_user(request)

router = APIRouter(prefix="/api/schemes", tags=["schemes"])

MGMT_ROLES = ("SuperAdmin", "Admin", "CEO", "Principal", "HeadTeacher", "Founder", "HeadOfAcademics")

def gen_id(prefix="sow_"):
    return f"{prefix}{uuid.uuid4().hex[:12]}"

def clean_ai_html(raw):
    """Strip code fences and full HTML doc wrappers, return only body content. Preserves math symbols."""
    text = raw.strip()
    text = re.sub(r'^```(?:html)?\s*\n?', '', text)
    text = re.sub(r'\n?```\s*$', '', text)
    # Strip full document wrapper if AI returns one
    body_match = re.search(r'<body[^>]*>(.*)</body>', text, re.DOTALL | re.IGNORECASE)
    if body_match:
        text = body_match.group(1).strip()
    else:
        text = re.sub(r'<!DOCTYPE[^>]*>', '', text, flags=re.IGNORECASE)
        text = re.sub(r'</?html[^>]*>', '', text, flags=re.IGNORECASE)
        text = re.sub(r'<head>.*?</head>', '', text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r'<style>.*?</style>', '', text, flags=re.DOTALL | re.IGNORECASE)
    return ensure_table_structure(text.strip())

def ensure_table_structure(html):
    """Ensure AI output has proper <table>, <thead>, <th> structure and convert markdown to HTML.
       Preserves mathematical formulas, subscripts, superscripts, and scientific notation."""
    if not html:
        return html
    # Convert markdown bold/italic to HTML — but protect math expressions first
    # Protect content inside <sub>, <sup>, <code>, <math>, <var> tags
    protected = {}
    counter = [0]
    def protect(m):
        key = f"\x00PROT{counter[0]}\x00"
        counter[0] += 1
        protected[key] = m.group(0)
        return key
    # Protect existing HTML tags that contain math/science content
    html = re.sub(r'<(sub|sup|code|math|var|pre|kbd)[^>]*>.*?</\1>', protect, html, flags=re.DOTALL | re.IGNORECASE)
    # Protect LaTeX-style expressions
    html = re.sub(r'\\\(.*?\\\)', protect, html, flags=re.DOTALL)
    html = re.sub(r'\\\[.*?\\\]', protect, html, flags=re.DOTALL)
    html = re.sub(r'\$\$.+?\$\$', protect, html, flags=re.DOTALL)
    html = re.sub(r'(?<!\$)\$(?!\$).+?\$(?!\$)', protect, html, flags=re.DOTALL)

    # Convert markdown bold/italic to HTML (only triple and double asterisks — single * is too ambiguous for math)
    html = re.sub(r'\*\*\*(.+?)\*\*\*', r'<strong><em>\1</em></strong>', html)
    html = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', html)
    html = re.sub(r'__(.+?)__', r'<strong>\1</strong>', html)

    # Restore protected content
    for key, val in protected.items():
        html = html.replace(key, val)

    has_table = bool(re.search(r'<table', html, re.IGNORECASE))
    has_tr = bool(re.search(r'<tr', html, re.IGNORECASE))
    has_td = bool(re.search(r'<td', html, re.IGNORECASE))
    # If we have <tr>/<td> but no <table>, wrap it
    if has_tr and has_td and not has_table:
        html = f"<table>\n<tbody>\n{html}\n</tbody>\n</table>"
    # If we have <table> but no <thead>, try to convert first row with bold/header content
    if has_table and not re.search(r'<thead', html, re.IGNORECASE):
        first_tr = re.search(r'<tr[^>]*>(.*?)</tr>', html, re.DOTALL | re.IGNORECASE)
        if first_tr:
            row_content = first_tr.group(1)
            cells = re.findall(r'<td[^>]*>(.*?)</td>', row_content, re.DOTALL | re.IGNORECASE)
            if cells and all(len(re.sub(r'<[^>]+>', '', c).strip()) < 40 for c in cells):
                new_header = row_content.replace('<td', '<th').replace('</td>', '</th>')
                header_row = f"<thead><tr>{new_header}</tr></thead>"
                html = html.replace(first_tr.group(0), '', 1)
                html = html.replace('<table>', f'<table>\n{header_row}', 1)
                html = re.sub(r'<table([^>]*)>\s*\n?\s*{}'.format(re.escape(header_row)), 
                    lambda m: f'<table{m.group(1)}>\n{header_row}', html)
    return html

STATUS_FLOW = {
    "draft": ["in_progress", "submitted"],
    "in_progress": ["draft", "submitted"],
    "submitted": ["under_review"],
    "under_review": ["approved", "returned"],
    "returned": ["in_progress", "resubmitted"],
    "resubmitted": ["under_review"],
    "approved": ["published", "archived"],
    "published": ["archived"],
    "archived": [],
}

STATUS_COLORS = {
    "draft": "gray", "in_progress": "blue", "submitted": "indigo",
    "under_review": "amber", "returned": "red", "resubmitted": "orange",
    "approved": "green", "published": "emerald", "archived": "slate"
}

async def log_scheme_action(scheme_id, user_id, user_name, action, details=""):
    await db.scheme_audit.insert_one({
        "audit_id": gen_id("saud_"),
        "scheme_id": scheme_id, "user_id": user_id, "user_name": user_name,
        "action": action, "details": details,
        "timestamp": datetime.now(timezone.utc).isoformat()
    })


# ─── STAFF DASHBOARD ───
@router.get("/my-dashboard")
async def my_dashboard(user: dict = Depends(get_current_user)):
    uid = user["user_id"]
    active_term, _ = await get_active_term()
    total = await db.schemes.count_documents({"teacher_id": uid})
    term_schemes = await db.schemes.count_documents({"teacher_id": uid, "term": active_term})
    drafts = await db.schemes.count_documents({"teacher_id": uid, "status": {"$in": ["draft", "in_progress"]}})
    submitted = await db.schemes.count_documents({"teacher_id": uid, "status": {"$in": ["submitted", "under_review", "resubmitted"]}})
    returned = await db.schemes.count_documents({"teacher_id": uid, "status": "returned"})
    approved = await db.schemes.count_documents({"teacher_id": uid, "status": {"$in": ["approved", "published"]}})
    recent_feedback = await db.scheme_comments.find(
        {"scheme_id": {"$in": [s["scheme_id"] async for s in db.schemes.find({"teacher_id": uid}, {"scheme_id": 1})]}, "is_management": True},
        {"_id": 0}
    ).sort("timestamp", -1).to_list(5)
    return {
        "total": total, "term_schemes": term_schemes, "drafts": drafts,
        "submitted": submitted, "returned": returned, "approved": approved,
        "active_term": active_term, "recent_feedback": recent_feedback
    }


# ─── MANAGEMENT DASHBOARD ───
@router.get("/management-dashboard")
async def management_dashboard(user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management access only")
    active_term, _ = await get_active_term()
    total = await db.schemes.count_documents({"status": {"$ne": "draft"}})
    pending = await db.schemes.count_documents({"status": {"$in": ["submitted", "resubmitted"]}})
    under_review = await db.schemes.count_documents({"status": "under_review"})
    approved_count = await db.schemes.count_documents({"status": {"$in": ["approved", "published"]}})
    returned_count = await db.schemes.count_documents({"status": "returned"})
    # Department stats
    pipeline = [
        {"$match": {"status": {"$ne": "draft"}}},
        {"$group": {"_id": "$department", "total": {"$sum": 1},
                    "approved": {"$sum": {"$cond": [{"$in": ["$status", ["approved", "published"]]}, 1, 0]}},
                    "pending": {"$sum": {"$cond": [{"$in": ["$status", ["submitted", "resubmitted", "under_review"]]}, 1, 0]}}}}
    ]
    dept_stats = await db.schemes.aggregate(pipeline).to_list(50)
    # Staff who haven't submitted
    all_teachers = await db.users.find(
        {"role": {"$nin": ["Student", "Parent"]}, "roles": {"$in": ["Teacher"]}},
        {"_id": 0, "user_id": 1, "name": 1, "department": 1, "teaching_subjects": 1}
    ).to_list(500)
    submitted_teachers = set()
    async for s in db.schemes.find({"term": active_term, "status": {"$nin": ["draft"]}}, {"teacher_id": 1}):
        submitted_teachers.add(s["teacher_id"])
    overdue_staff = [t for t in all_teachers if t["user_id"] not in submitted_teachers]
    return {
        "total": total, "pending": pending, "under_review": under_review,
        "approved": approved_count, "returned": returned_count,
        "department_stats": dept_stats, "overdue_staff": overdue_staff,
        "active_term": active_term
    }


# ─── LIST SCHEMES ───
@router.get("/")
async def list_schemes(
    user: dict = Depends(get_current_user),
    status: str = Query(None), department: str = Query(None),
    subject: str = Query(None), teacher_id: str = Query(None),
    term: str = Query(None), year_group: str = Query(None),
):
    is_mgmt = user.get("role") in MGMT_ROLES
    query = {}
    if not is_mgmt:
        query["teacher_id"] = user["user_id"]
    else:
        if status == "all_submitted":
            query["status"] = {"$ne": "draft"}
        elif status:
            query["status"] = status
        if teacher_id:
            query["teacher_id"] = teacher_id
    if department:
        query["department"] = department
    if subject:
        query["subject_id"] = subject
    if term:
        query["term"] = term
    if year_group:
        query["year_group"] = year_group
    if not is_mgmt and status:
        query["status"] = status
    schemes = await db.schemes.find(query, {"_id": 0}).sort("updated_at", -1).to_list(500)
    # Filter by approval chain visibility for management
    if is_mgmt:
        schemes = [s for s in schemes if s.get("teacher_id") == user["user_id"] or can_user_see_lesson(user, s.get("approval_chain", []))]
    return schemes


# ─── CREATE SCHEME (from scratch) ───
@router.post("/")
async def create_scheme(data: dict, user: dict = Depends(get_current_user)):
    active_term, academic_year = await get_active_term()
    scheme_id = gen_id()
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "scheme_id": scheme_id,
        "teacher_id": user["user_id"],
        "teacher_name": user.get("name", ""),
        "department": data.get("department", user.get("department", "")),
        "subject_id": data.get("subject_id", ""),
        "subject_name": data.get("subject_name", ""),
        "year_group": data.get("year_group", ""),
        "class_name": data.get("class_name", ""),
        "curriculum": data.get("curriculum", "Cambridge"),
        "term": data.get("term", active_term),
        "academic_year": data.get("academic_year", academic_year),
        "title": data.get("title", f"Scheme of Work - {data.get('subject_name', '')} {data.get('year_group', '')}"),
        "weeks": data.get("weeks", []),
        "rich_content": data.get("rich_content", ""),
        "original_file_id": None,
        "ai_processed": False,
        "ai_version": None,
        "status": "draft",
        "version": 1,
        "versions": [],
        "reviewer_id": None,
        "reviewer_name": None,
        "approval_notes": "",
        "resubmission_count": 0,
        "created_at": now,
        "updated_at": now,
        "submitted_at": None,
        "approved_at": None,
        "published_at": None,
    }
    await db.schemes.insert_one(doc)
    await log_scheme_action(scheme_id, user["user_id"], user["name"], "created", "Scheme created from scratch")
    # Exclude _id from response (MongoDB adds it during insert)
    doc.pop("_id", None)
    return doc


# ─── BATCH PDF EXPORT (Management) ───
@router.get("/batch-export-pdf")
async def batch_export_pdf(
    request: Request,
    token: str = Query(None),
    term: str = Query(None),
    department: str = Query(None),
    year_group: str = Query(None),
    status: str = Query(None),
):
    user = await _get_user_from_token_or_header(request, token)
    if not any(r in user.get("roles", [user.get("role", "")]) for r in MGMT_ROLES):
        raise HTTPException(status_code=403, detail="Management only")

    query = {}
    if term:
        query["term"] = term
    if department:
        query["department"] = department
    if year_group:
        query["year_group"] = year_group
    if status:
        query["status"] = status
    else:
        query["status"] = {"$in": ["submitted", "under_review", "approved", "published"]}

    schemes = await db.schemes.find(query, {"_id": 0}).sort("department", 1).to_list(200)
    schemes = [s for s in schemes if s.get("rich_content")]
    if not schemes:
        raise HTTPException(status_code=404, detail="No schemes found matching filters")

    logo_path = os.path.join(os.path.dirname(__file__), "..", "main_logo.png")
    logo_b64 = ""
    if os.path.exists(logo_path):
        with open(logo_path, "rb") as f:
            logo_b64 = base64.b64encode(f.read()).decode()

    pages_html = []
    for i, scheme in enumerate(schemes):
        page_break = 'page-break-before: always;' if i > 0 else ''
        pages_html.append(f"""
<div style="{page_break}">
<div class="header">
  {"<img src='data:image/png;base64," + logo_b64 + "' />" if logo_b64 else ""}
  <h1>The RightWay Centre of Learning</h1>
  <h2>{scheme.get('title', 'Scheme of Work')}</h2>
</div>
<div class="meta">
  <div>
    <span>Subject: <strong>{scheme.get('subject_name', '')}</strong></span>
    <span>Year Group: <strong>{scheme.get('year_group', '')}</strong></span>
    <span>Term: <strong>{scheme.get('term', '')}</strong></span>
    <span>Department: <strong>{scheme.get('department', '')}</strong></span>
  </div>
  <div>
    <span>Teacher: <strong>{scheme.get('teacher_name', '')}</strong></span>
    <span>Status: <strong>{scheme.get('status', '').replace('_', ' ').title()}</strong></span>
    <span>Curriculum: <strong>{scheme.get('curriculum', 'Cambridge')}</strong></span>
  </div>
</div>
{scheme.get('rich_content', '')}
</div>""")

    html_doc = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
@page {{ size: A4 landscape; margin: 15mm 12mm; }}
body {{ font-family: Arial, Helvetica, sans-serif; font-size: 11px; color: #1a1a1a; line-height: 1.4; }}
.header {{ text-align: center; margin-bottom: 14px; padding-bottom: 10px; border-bottom: 3px solid #1B2A4A; }}
.header img {{ width: 180px; height: auto; margin-bottom: 6px; }}
.header h1 {{ font-size: 16px; color: #1B2A4A; margin: 4px 0 2px; text-transform: uppercase; letter-spacing: 1px; }}
.header h2 {{ font-size: 13px; color: #2E6CC4; margin: 2px 0; font-weight: 600; }}
.meta {{ display: flex; justify-content: space-between; margin-bottom: 10px; font-size: 10px; color: #555; }}
.meta span {{ margin-right: 16px; }}
.meta strong {{ color: #1B2A4A; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 6px; font-size: 10px; }}
thead th {{ background: #1B2A4A; color: white; padding: 7px 6px; text-align: left; font-weight: 600; font-size: 9px; text-transform: uppercase; letter-spacing: 0.4px; border: 1px solid #1B2A4A; }}
tbody td {{ padding: 6px; border: 1px solid #d1d5db; vertical-align: top; line-height: 1.35; }}
tbody tr:nth-child(even) td {{ background: #f8fafc; }}
ul, ol {{ margin: 2px 0; padding-left: 14px; }}
li {{ margin-bottom: 1px; }}
strong {{ font-weight: 700; }}
sub {{ font-size: 0.75em; vertical-align: sub; }}
sup {{ font-size: 0.75em; vertical-align: super; }}
var {{ font-style: italic; }}
.footer {{ text-align: center; font-size: 8px; color: #999; margin-top: 10px; padding-top: 6px; border-top: 1px solid #e5e7eb; }}
</style></head><body>
{''.join(pages_html)}
<div class="footer">TCIS Batch Report — {len(schemes)} Scheme(s) — Generated {datetime.now().strftime('%d %B %Y')} — Confidential</div>
</body></html>"""

    try:
        from weasyprint import HTML
        pdf_bytes = HTML(string=html_doc).write_pdf()
        dept_label = department.replace(' ', '_') if department else 'All'
        filename = f"TCIS_Schemes_{dept_label}_{term or 'AllTerms'}.pdf"
        return Response(content=pdf_bytes, media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={filename}"})
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Batch PDF failed: {str(e)}")


# ─── MY PENDING REVIEWS (Schemes) — must be before /{scheme_id} ───
@router.get("/my-pending-reviews")
async def my_pending_scheme_reviews(user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    all_pending = await db.schemes.find(
        {"status": {"$in": ["submitted", "under_review", "resubmitted"]}, "approval_chain": {"$exists": True}},
        {"_id": 0},
    ).sort("submitted_at", -1).to_list(200)
    my_items = []
    for item in all_pending:
        chain = item.get("approval_chain", [])
        idx = get_current_step_index(chain)
        if idx >= 0 and can_user_approve_step(user, chain[idx], chain):
            my_items.append(item)
    return my_items


# ─── GET SINGLE SCHEME ───
@router.get("/{scheme_id}")
async def get_scheme(scheme_id: str, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Scheme not found")
    comments = await db.scheme_comments.find({"scheme_id": scheme_id}, {"_id": 0}).sort("timestamp", 1).to_list(200)
    audit = await db.scheme_audit.find({"scheme_id": scheme_id}, {"_id": 0}).sort("timestamp", -1).to_list(100)
    return {**scheme, "comments": comments, "audit_trail": audit}


# ─── UPDATE SCHEME (staff editing) ───
@router.put("/{scheme_id}")
async def update_scheme(scheme_id: str, data: dict, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    if scheme["teacher_id"] != user["user_id"] and user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Not your scheme")
    if scheme["status"] in ("approved", "published", "archived"):
        raise HTTPException(status_code=400, detail="Cannot edit approved/published scheme")
    safe = {"title", "department", "subject_id", "subject_name", "year_group", "class_name",
            "curriculum", "term", "academic_year", "weeks", "rich_content", "original_file_id"}
    update = {k: v for k, v in data.items() if k in safe}
    update["updated_at"] = datetime.now(timezone.utc).isoformat()
    if scheme["status"] == "draft":
        update["status"] = "in_progress"
    await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": update})
    await log_scheme_action(scheme_id, user["user_id"], user["name"], "edited", "Content updated")
    updated = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    return updated


# ─── SUBMIT TO MANAGEMENT ───
@router.post("/{scheme_id}/submit")
async def submit_scheme(scheme_id: str, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    if scheme["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your scheme")
    if scheme["status"] not in ("draft", "in_progress", "returned"):
        raise HTTPException(status_code=400, detail=f"Cannot submit from '{scheme['status']}' status")
    now = datetime.now(timezone.utc).isoformat()
    new_status = "resubmitted" if scheme["status"] == "returned" else "submitted"
    chain = build_approval_chain(user)
    if is_chain_complete(chain):
        new_status = "approved"
    current_idx = get_current_step_index(chain)
    update = {
        "status": new_status, "submitted_at": now, "updated_at": now,
        "approval_chain": chain,
        "current_approval_step": chain[current_idx]["step"] if current_idx >= 0 else 0,
        "current_approver_label": get_current_step_label(chain),
    }
    if new_status == "approved":
        update["approved_at"] = now
    if scheme.get("status") == "returned":
        update["resubmission_count"] = scheme.get("resubmission_count", 0) + 1
    version_snap = {"version": scheme.get("version", 1), "weeks": scheme.get("weeks", []),
                    "rich_content": scheme.get("rich_content", ""), "saved_at": now}
    await db.schemes.update_one({"scheme_id": scheme_id}, {
        "$set": update, "$push": {"versions": version_snap}, "$inc": {"version": 1}
    })
    await log_scheme_action(scheme_id, user["user_id"], user["name"], new_status,
                            f"Scheme {new_status} — {get_current_step_label(chain)}")
    return {"message": f"Scheme {new_status} successfully", "status": new_status, "current_approver": get_current_step_label(chain)}


# ─── MANAGEMENT: REVIEW ACTION (Hierarchical) ───
@router.post("/{scheme_id}/review")
async def review_scheme(scheme_id: str, data: dict, user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    action = data.get("action")
    comment_text = data.get("comment", "")
    now = datetime.now(timezone.utc).isoformat()
    chain = scheme.get("approval_chain", [])

    if action in ("approve", "return") and chain:
        current_idx = get_current_step_index(chain)
        if current_idx == -1:
            raise HTTPException(status_code=400, detail="All steps already completed")
        current_step = chain[current_idx]
        if not can_user_approve_step(user, current_step, chain):
            raise HTTPException(status_code=403,
                                detail=f"This item is currently with {current_step['label']}. You cannot act on it at this stage.")

        if action == "approve":
            chain[current_idx]["status"] = "approved"
            chain[current_idx]["user_id"] = user["user_id"]
            chain[current_idx]["user_name"] = user.get("name", "")
            chain[current_idx]["acted_at"] = now
            chain[current_idx]["comment"] = comment_text
            if is_chain_complete(chain):
                await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
                    "status": "approved", "approved_at": now, "approval_notes": comment_text,
                    "approval_chain": chain, "current_approval_step": 0,
                    "current_approver_label": "Fully Approved",
                    "reviewer_id": user["user_id"], "reviewer_name": user.get("name", ""), "updated_at": now,
                }})
                await log_scheme_action(scheme_id, user["user_id"], user["name"], "approved",
                                        f"Final approval by {current_step['label']}. {comment_text}")
            else:
                next_idx = get_current_step_index(chain)
                next_label = get_current_step_label(chain)
                await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
                    "status": "under_review", "approval_chain": chain,
                    "current_approval_step": chain[next_idx]["step"],
                    "current_approver_label": next_label,
                    "reviewer_id": user["user_id"], "reviewer_name": user.get("name", ""), "updated_at": now,
                }})
                await log_scheme_action(scheme_id, user["user_id"], user["name"], "step_approved",
                                        f"Approved by {current_step['label']} → Now {next_label}. {comment_text}")
        elif action == "return":
            chain[current_idx]["status"] = "returned"
            chain[current_idx]["user_id"] = user["user_id"]
            chain[current_idx]["user_name"] = user.get("name", "")
            chain[current_idx]["acted_at"] = now
            chain[current_idx]["comment"] = comment_text
            await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
                "status": "returned", "approval_chain": chain,
                "current_approval_step": current_step["step"],
                "current_approver_label": f"Returned by {current_step['label']}",
                "reviewer_id": user["user_id"], "reviewer_name": user.get("name", ""), "updated_at": now,
            }})
            await log_scheme_action(scheme_id, user["user_id"], user["name"], "returned",
                                    f"Returned by {current_step['label']}. {comment_text}")
    elif action == "start_review":
        await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
            "status": "under_review", "reviewer_id": user["user_id"],
            "reviewer_name": user.get("name", ""), "updated_at": now
        }})
        await log_scheme_action(scheme_id, user["user_id"], user["name"], "under_review", "Review started")
    elif action == "publish":
        await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
            "status": "published", "published_at": now, "updated_at": now
        }})
        await log_scheme_action(scheme_id, user["user_id"], user["name"], "published", "Scheme published")
    elif action == "archive":
        await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
            "status": "archived", "updated_at": now
        }})
        await log_scheme_action(scheme_id, user["user_id"], user["name"], "archived", "Scheme archived")
    if comment_text:
        await db.scheme_comments.insert_one({
            "comment_id": gen_id("scmt_"), "scheme_id": scheme_id,
            "user_id": user["user_id"], "user_name": user.get("name", ""),
            "text": comment_text, "is_management": True, "action": action,
            "timestamp": now
        })
    return {"message": f"Action '{action}' completed", "status": action}




# ─── ADD COMMENT ───
@router.post("/{scheme_id}/comments")
async def add_comment(scheme_id: str, data: dict, user: dict = Depends(get_current_user)):
    is_mgmt = user.get("role") in MGMT_ROLES
    now = datetime.now(timezone.utc).isoformat()
    await db.scheme_comments.insert_one({
        "comment_id": gen_id("scmt_"), "scheme_id": scheme_id,
        "user_id": user["user_id"], "user_name": user.get("name", ""),
        "text": data.get("text", ""), "is_management": is_mgmt,
        "section": data.get("section", ""), "timestamp": now
    })
    await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {"updated_at": now}})
    return {"message": "Comment added"}


# ─── COMPLIANCE TRACKER ───
@router.get("/management/compliance")
async def compliance_tracker(user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    active_term, _ = await get_active_term()
    teachers = await db.users.find(
        {"roles": {"$in": ["Teacher"]}},
        {"_id": 0, "user_id": 1, "name": 1, "department": 1, "teaching_subjects": 1}
    ).to_list(500)
    schemes = await db.schemes.find({"term": active_term}, {"_id": 0, "teacher_id": 1, "subject_id": 1, "status": 1}).to_list(1000)
    scheme_map = {}
    for s in schemes:
        key = s["teacher_id"]
        if key not in scheme_map:
            scheme_map[key] = []
        scheme_map[key].append(s)
    result = []
    for t in teachers:
        t_schemes = scheme_map.get(t["user_id"], [])
        subjects = t.get("teaching_subjects", [])
        required = max(len(subjects), 1)
        submitted_count = sum(1 for s in t_schemes if s["status"] not in ("draft",))
        approved_count = sum(1 for s in t_schemes if s["status"] in ("approved", "published"))
        result.append({
            "user_id": t["user_id"], "name": t["name"],
            "department": t.get("department", ""), "subjects": subjects,
            "required": required, "submitted": submitted_count,
            "approved": approved_count, "pending": submitted_count - approved_count,
            "missing": max(required - submitted_count, 0),
            "rate": round(submitted_count / required * 100) if required > 0 else 0
        })
    result.sort(key=lambda x: x["rate"])
    return result


# ─── AI IMPROVE SCHEME ───
@router.post("/{scheme_id}/ai-improve")
async def ai_improve_scheme(scheme_id: str, data: dict, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    if scheme["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your scheme")
    prompt_type = data.get("prompt_type", "improve")
    custom_prompt = data.get("custom_prompt", "")
    content = scheme.get("rich_content", "")
    weeks_data = scheme.get("weeks", [])
    weeks_text = ""
    for w in weeks_data:
        weeks_text += f"\nWeek {w.get('week_number', '?')}: Topic: {w.get('topic', '')}, Objectives: {w.get('objectives', '')}, Activities: {w.get('activities', '')}, Resources: {w.get('resources', '')}, Assessment: {w.get('assessment', '')}"
    full_content = f"Subject: {scheme.get('subject_name', '')}\nYear Group: {scheme.get('year_group', '')}\nTerm: {scheme.get('term', '')}\nClass: {scheme.get('class_name', '')}\n\n{content}\n{weeks_text}"
    MATH_RULE = """\n\nMATH & SCIENCE NOTATION: You MUST preserve ALL mathematical formulas, symbols, and scientific notation exactly. Use HTML tags:
- Subscripts: <sub>2</sub> for H<sub>2</sub>O, CO<sub>2</sub>, x<sub>n</sub>
- Superscripts: <sup>2</sup> for x<sup>2</sup>, m<sup>3</sup>, 10<sup>8</sup>
- Keep Unicode symbols as-is: √, π, θ, α, β, γ, δ, Σ, ∫, ∞, ≤, ≥, ≠, ±, ÷, ×, °, Ω, μ, Δ
- Fractions: use <sup>a</sup>/<sub>b</sub> notation
- Chemical equations: H<sub>2</sub>SO<sub>4</sub>, NaOH, Fe<sub>2</sub>O<sub>3</sub>
- Physics: F = ma, E = mc<sup>2</sup>, v = u + at
NEVER replace, simplify, or strip these symbols. They are essential academic content."""
    TABLE_RULE = "\n\nCRITICAL: If the content contains an HTML table, you MUST preserve the exact <table>, <thead>, <tr>, <th>, <tbody>, <td> structure. Return the COMPLETE table with header row using <th> inside <thead>. No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags." + MATH_RULE
    prompts = {
        "improve": "You are an expert academic planning assistant for TCIS (The RightWay Centre of Learning). Improve this scheme of work professionally. Fix formatting, strengthen learning objectives with measurable verbs, improve clarity, ensure alignment between objectives, activities, and assessments." + TABLE_RULE,
        "standardise": "Rewrite this scheme of work in TCIS standard academic planning format. Structure it clearly by weeks with topic, sub-topic, learning objectives, teaching activities, resources, assessment, and homework." + TABLE_RULE,
        "expand": "Expand weak or incomplete sections of this scheme of work. Add more detail to learning objectives, suggest specific teaching activities, add assessment methods where missing, suggest relevant resources." + TABLE_RULE,
        "objectives": "Rewrite all learning objectives in this scheme using Bloom's taxonomy measurable verbs (analyze, evaluate, create, apply, etc.). Make each objective specific, measurable, achievable, relevant, and time-bound." + TABLE_RULE,
        "assessment": "Add or improve assessment strategies throughout this scheme of work. Suggest formative and summative assessment methods appropriate for each topic." + TABLE_RULE,
        "gaps": "Analyze this scheme of work and identify: missing weeks, weak objectives, poor sequencing, missing assessments, curriculum gaps, formatting issues. Return a structured HTML report with specific recommendations." + TABLE_RULE,
        "digitalize": "Convert this raw content into a clean, structured digital scheme of work in HTML table format with columns: Week, Topic, Sub-Topic, Learning Objectives, Activities, Resources, Assessment, Homework. Preserve all original content but organize it professionally." + TABLE_RULE,
        "custom": (custom_prompt or "Improve this scheme of work.") + TABLE_RULE
    }
    system_prompt = prompts.get(prompt_type, prompts["improve"])
    try:
        from services.ai_service import llm_chat
        response = await llm_chat(system_prompt, f"Here is the scheme of work content:\n\n{full_content}", provider="gemini")
        ai_content = clean_ai_html(response)
        now = datetime.now(timezone.utc).isoformat()
        await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
            "ai_version": ai_content, "ai_processed": True,
            "rich_content": ai_content, "updated_at": now
        }})
        await log_scheme_action(scheme_id, user["user_id"], user["name"], "ai_improved", f"AI {prompt_type}")
        return {"ai_content": ai_content, "prompt_type": prompt_type}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI processing failed: {str(e)}")



# ─── AI ANALYZE SCHEME (find issues) ───
@router.post("/{scheme_id}/ai-analyze")
async def ai_analyze_scheme(scheme_id: str, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    if scheme["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your scheme")
    content = scheme.get("rich_content", "")
    if not content or len(content) < 30:
        return {"issues": []}
    try:
        from services.ai_service import llm_chat
        system = """You are a strict academic quality reviewer for TCIS (The RightWay Centre of Learning).
Analyze the scheme of work and return ONLY a valid JSON array of issues. Each issue object must have:
- "id": a short unique slug (e.g. "weak-objectives-w3")
- "category": one of "Formatting", "Objectives", "Assessment", "Resources", "Curriculum Gap", "Sequencing", "Alignment"
- "severity": one of "high", "medium", "low"
- "week": the affected week number or "all" if it applies globally
- "title": a short 5-10 word summary of the issue
- "description": 1-2 sentences explaining the issue and what should be fixed
Return ONLY the JSON array, no markdown, no explanation, no code fences."""
        response = await llm_chat(system, f"Subject: {scheme.get('subject_name','')}\nYear: {scheme.get('year_group','')}\nTerm: {scheme.get('term','')}\n\nScheme content:\n{content}", provider="gemini")
        import json
        text = response.strip()
        text = re.sub(r'^```(?:json)?\s*\n?', '', text)
        text = re.sub(r'\n?```\s*$', '', text)
        issues = json.loads(text)
        if not isinstance(issues, list):
            issues = []
        # Store issues on the scheme
        await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {"ai_issues": issues}})
        return {"issues": issues}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Analysis failed: {str(e)}")


# ─── AI FIX SPECIFIC ISSUE ───
@router.post("/{scheme_id}/ai-fix-issue")
async def ai_fix_issue(scheme_id: str, data: dict, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    if scheme["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your scheme")
    issue_title = data.get("issue_title", "")
    issue_description = data.get("issue_description", "")
    issue_id = data.get("issue_id", "")
    content = scheme.get("rich_content", "")
    if not content:
        raise HTTPException(status_code=400, detail="No content to fix")
    try:
        from services.ai_service import llm_chat
        system = f"""You are an expert academic planning assistant for TCIS school.
You must fix ONLY this specific issue in the scheme of work: "{issue_title}" - {issue_description}

CRITICAL RULES:
- Keep the EXACT same HTML table structure: <table>, <thead>, <tr>, <th>, <tbody>, <td> tags
- Do NOT remove or change the table wrapper, header row, or column structure
- Only modify the specific cell content related to the issue
- Keep all other content exactly the same
- Return the COMPLETE HTML including the full <table> with <thead> and <tbody>
- No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags"""
        response = await llm_chat(system, f"Here is the full scheme of work. Fix ONLY the issue described above:\n\n{content}", provider="gemini")
        ai_content = clean_ai_html(response)
        now = datetime.now(timezone.utc).isoformat()
        # Remove the fixed issue from stored issues
        existing_issues = scheme.get("ai_issues", [])
        remaining_issues = [i for i in existing_issues if i.get("id") != issue_id]
        await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
            "rich_content": ai_content, "ai_version": ai_content,
            "ai_processed": True, "ai_issues": remaining_issues, "updated_at": now
        }})
        await log_scheme_action(scheme_id, user["user_id"], user["name"], "ai_fix_issue", f"Fixed: {issue_title}")
        return {"ai_content": ai_content, "remaining_issues": remaining_issues}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Fix failed: {str(e)}")

# ─── AI FIX ALL ISSUES ───
@router.post("/{scheme_id}/ai-fix-all")
async def ai_fix_all_issues(scheme_id: str, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    if scheme["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your scheme")
    issues = scheme.get("ai_issues", [])
    content = scheme.get("rich_content", "")
    if not content or not issues:
        raise HTTPException(status_code=400, detail="No content or no issues to fix")
    issues_text = "\n".join([f"- [{i.get('severity','medium').upper()}] {i.get('title','')}: {i.get('description','')}" for i in issues])
    try:
        from services.ai_service import llm_chat
        system = f"""You are an expert academic planning assistant for TCIS school.
Fix ALL of the following issues in the scheme of work:
{issues_text}

CRITICAL RULES:
- Keep the EXACT same HTML table structure: <table>, <thead>, <tr>, <th>, <tbody>, <td> tags
- Do NOT remove or change the table wrapper, header row, or column structure
- Only modify the specific cell content related to each issue
- Return the COMPLETE HTML including the full <table> with <thead> and <tbody>
- The header row must use <th> tags inside <thead>
- No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags"""
        response = await llm_chat(system, f"Here is the full scheme of work:\n\n{content}", provider="gemini")
        ai_content = clean_ai_html(response)
        now = datetime.now(timezone.utc).isoformat()
        await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
            "rich_content": ai_content, "ai_version": ai_content,
            "ai_processed": True, "ai_issues": [], "updated_at": now
        }})
        await log_scheme_action(scheme_id, user["user_id"], user["name"], "ai_fix_all", f"Fixed {len(issues)} issues")
        return {"ai_content": ai_content, "fixed_count": len(issues)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Fix all failed: {str(e)}")



# ─── PDF EXPORT ───
@router.get("/{scheme_id}/export-pdf")
async def export_scheme_pdf(scheme_id: str, request: Request, token: str = Query(None), download: int = Query(0)):
    await _get_user_from_token_or_header(request, token)
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    content = scheme.get("rich_content", "")
    if not content:
        content = "<p style='color:#888; text-align:center; padding:40px;'>This scheme of work has no content yet. Add content in the editor to generate a full PDF.</p>"

    # Encode logo as base64 data URI
    logo_path = os.path.join(os.path.dirname(__file__), "..", "main_logo.png")
    logo_b64 = ""
    if os.path.exists(logo_path):
        with open(logo_path, "rb") as f:
            logo_b64 = base64.b64encode(f.read()).decode()

    html_doc = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
@page {{ size: A4 landscape; margin: 15mm 12mm; }}
body {{ font-family: Arial, Helvetica, sans-serif; font-size: 11px; color: #1a1a1a; line-height: 1.4; }}
.header {{ text-align: center; margin-bottom: 14px; padding-bottom: 10px; border-bottom: 3px solid #1B2A4A; }}
.header img {{ width: 180px; height: auto; margin-bottom: 6px; }}
.header h1 {{ font-size: 16px; color: #1B2A4A; margin: 4px 0 2px; text-transform: uppercase; letter-spacing: 1px; }}
.header h2 {{ font-size: 13px; color: #2E6CC4; margin: 2px 0; font-weight: 600; }}
.meta {{ display: flex; justify-content: space-between; margin-bottom: 10px; font-size: 10px; color: #555; }}
.meta span {{ margin-right: 16px; }}
.meta strong {{ color: #1B2A4A; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 6px; font-size: 10px; }}
thead th {{ background: #1B2A4A; color: white; padding: 7px 6px; text-align: left; font-weight: 600; font-size: 9px; text-transform: uppercase; letter-spacing: 0.4px; border: 1px solid #1B2A4A; }}
tbody td {{ padding: 6px; border: 1px solid #d1d5db; vertical-align: top; line-height: 1.35; }}
tbody tr:nth-child(even) td {{ background: #f8fafc; }}
ul, ol {{ margin: 2px 0; padding-left: 14px; }}
li {{ margin-bottom: 1px; }}
strong {{ font-weight: 700; }}
sub {{ font-size: 0.75em; vertical-align: sub; }}
sup {{ font-size: 0.75em; vertical-align: super; }}
var {{ font-style: italic; }}
.footer {{ text-align: center; font-size: 8px; color: #999; margin-top: 10px; padding-top: 6px; border-top: 1px solid #e5e7eb; }}
</style>
</head>
<body>
<div class="header">
  {"<img src='data:image/png;base64," + logo_b64 + "' />" if logo_b64 else ""}
  <h1>The RightWay Centre of Learning</h1>
  <h2>{scheme.get('title', 'Scheme of Work')}</h2>
</div>
<div class="meta">
  <div>
    <span>Subject: <strong>{scheme.get('subject_name', '')}</strong></span>
    <span>Year Group: <strong>{scheme.get('year_group', '')}</strong></span>
    <span>Term: <strong>{scheme.get('term', '')}</strong></span>
    <span>Department: <strong>{scheme.get('department', '')}</strong></span>
  </div>
  <div>
    <span>Teacher: <strong>{scheme.get('teacher_name', '')}</strong></span>
    <span>Curriculum: <strong>{scheme.get('curriculum', 'Cambridge')}</strong></span>
    <span>Academic Year: <strong>{scheme.get('academic_year', '')}</strong></span>
  </div>
</div>
{content}
<div class="footer">TCIS Scheme of Work — Generated {datetime.now().strftime('%d %B %Y')} — Confidential</div>
</body>
</html>"""

    try:
        from weasyprint import HTML
        pdf_bytes = HTML(string=html_doc).write_pdf()
        filename = f"SOW_{scheme.get('subject_name','').replace(' ','_')}_{scheme.get('year_group','').replace(' ','_')}.pdf"
        disposition = "attachment" if download else "inline"
        return Response(content=pdf_bytes, media_type="application/pdf",
            headers={"Content-Disposition": f"{disposition}; filename={filename}"})
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF generation failed: {str(e)}")


# ─── AI EVALUATE SCHEME (Management grading) ───
@router.post("/{scheme_id}/ai-evaluate")
async def ai_evaluate_scheme(scheme_id: str, user: dict = Depends(get_current_user)):
    if not any(r in user.get("roles", [user.get("role", "")]) for r in MGMT_ROLES):
        raise HTTPException(status_code=403, detail="Management only")
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    content = scheme.get("rich_content", "")
    if not content or len(content) < 30:
        raise HTTPException(status_code=400, detail="Not enough content to evaluate")
    try:
        from services.ai_service import llm_chat
        system = """You are a senior academic quality evaluator for TCIS (The RightWay Centre of Learning).
Grade this scheme of work and return ONLY a valid JSON object with this exact structure:
{
  "overall_score": <number 0-100>,
  "grade": "<A+/A/B+/B/C+/C/D/F>",
  "summary": "<2-3 sentence overall assessment>",
  "categories": [
    {"name": "Learning Objectives", "score": <0-100>, "feedback": "<specific feedback>"},
    {"name": "Teaching Activities", "score": <0-100>, "feedback": "<specific feedback>"},
    {"name": "Assessment Strategies", "score": <0-100>, "feedback": "<specific feedback>"},
    {"name": "Resource Provision", "score": <0-100>, "feedback": "<specific feedback>"},
    {"name": "Curriculum Alignment", "score": <0-100>, "feedback": "<specific feedback>"},
    {"name": "Sequencing & Pacing", "score": <0-100>, "feedback": "<specific feedback>"},
    {"name": "Differentiation", "score": <0-100>, "feedback": "<specific feedback>"},
    {"name": "Completeness", "score": <0-100>, "feedback": "<specific feedback>"}
  ],
  "errors": [
    {"week": "<week number or 'General'>", "issue": "<specific error description>", "severity": "<high/medium/low>", "suggestion": "<how to fix>"}
  ],
  "strengths": ["<strength 1>", "<strength 2>"],
  "recommendation": "<approve/return/needs_major_revision>"
}
Be strict and thorough. Only give 90+ for truly exceptional work. Return ONLY the JSON, no markdown, no code fences."""
        response = await llm_chat(system, f"Subject: {scheme.get('subject_name','')}\nYear: {scheme.get('year_group','')}\nTerm: {scheme.get('term','')}\nTeacher: {scheme.get('teacher_name','')}\n\nScheme content:\n{content}", provider="gemini")
        import json
        text = response.strip()
        text = re.sub(r'^```(?:json)?\s*\n?', '', text)
        text = re.sub(r'\n?```\s*$', '', text)
        evaluation = json.loads(text)
        # Store evaluation on the scheme
        now = datetime.now(timezone.utc).isoformat()
        await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
            "evaluation": evaluation, "evaluated_at": now, "evaluated_by": user["user_id"]
        }})
        await log_scheme_action(scheme_id, user["user_id"], user["name"], "evaluated", f"Score: {evaluation.get('overall_score', '?')}/100")
        return evaluation
    except json.JSONDecodeError:
        raise HTTPException(status_code=500, detail="AI returned invalid evaluation format")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Evaluation failed: {str(e)}")


# ─── AI PROCESS UPLOADED FILE ───
@router.post("/{scheme_id}/ai-process-file")
async def ai_process_file(scheme_id: str, data: dict, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    file_id = data.get("file_id") or scheme.get("original_file_id")
    if not file_id:
        raise HTTPException(status_code=400, detail="No file to process")
    # Get file content
    from storage import get_object
    import tempfile
    tmp_path = None
    try:
        file_meta = await db.file_assets.find_one({"file_id": file_id}, {"_id": 0})
        if not file_meta:
            raise HTTPException(status_code=404, detail="File record not found in database")
        storage_path = file_meta.get("storage_path", "")
        if not storage_path:
            raise HTTPException(status_code=400, detail="File has no storage path")
        file_bytes, content_type = get_object(storage_path)
        mime = file_meta.get("content_type", content_type or "application/pdf")
        ext = storage_path.rsplit('.', 1)[-1].lower() if '.' in storage_path else 'pdf'
        # Save to temp file
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=f'.{ext}')
        tmp.write(file_bytes)
        tmp.close()
        tmp_path = tmp.name
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"File read failed: {str(e)}")

    base_prompt = f"""Extract and digitalize this uploaded scheme of work document into a clean, professional HTML format.
Structure it as a table with: Week Number, Topic, Sub-Topic, Learning Objectives, Teaching Activities, Resources, Assessment Strategies, Homework.
If the document is messy or handwritten, do your best to interpret and organize the content.
Also identify: formatting issues, weak objectives, missing assessments, gaps, and areas for improvement.
At the end, add an 'AI Recommendations' section listing specific improvements.
Return ONLY inner HTML body content (tables, headings, paragraphs). Do NOT include <!DOCTYPE>, <html>, <head>, <style>, or <body> tags.

MATH & SCIENCE NOTATION: Preserve ALL mathematical formulas, symbols, and scientific notation exactly. Use HTML tags:
- Subscripts: <sub>2</sub> for H<sub>2</sub>O, CO<sub>2</sub>
- Superscripts: <sup>2</sup> for x<sup>2</sup>, m<sup>3</sup>
- Keep Unicode symbols: √, π, θ, α, β, Σ, ∫, ∞, ≤, ≥, ≠, ±, ÷, ×, °, Ω, μ, Δ
- Chemical equations: H<sub>2</sub>SO<sub>4</sub>, NaOH
- Physics: F = ma, E = mc<sup>2</sup>
NEVER simplify or strip mathematical/scientific symbols.

Subject: {scheme.get('subject_name', 'Unknown')}
Year Group: {scheme.get('year_group', '')}
Term: {scheme.get('term', '')}"""

    try:
        from services.ai_service import llm_chat, llm_chat_with_file
        system_for_file = "You are an expert academic planning assistant for The RightWay Centre of Learning."

        # For unsupported MIME types (docx, doc, etc.), extract text and send as text
        is_pdf = ext == 'pdf' or 'pdf' in mime
        is_image = mime.startswith('image/')
        use_file_upload = is_pdf or is_image  # Gemini supports PDF and images natively

        if use_file_upload:
            response = await llm_chat_with_file(system_for_file, base_prompt, file_path=tmp_path, mime_type=mime)
        else:
            # Extract text from docx/doc files
            extracted_text = ""
            if ext in ('docx', 'doc'):
                try:
                    from docx import Document
                    doc = Document(tmp_path)
                    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                    # Also extract text from tables
                    for table in doc.tables:
                        for row in table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells)
                            if row_text.strip():
                                paragraphs.append(row_text)
                    extracted_text = "\n".join(paragraphs)
                except Exception:
                    extracted_text = file_bytes.decode('utf-8', errors='ignore')
            else:
                extracted_text = file_bytes.decode('utf-8', errors='ignore')

            if not extracted_text.strip():
                extracted_text = "(Document appears empty or could not be read)"

            response = await llm_chat(system_for_file, f"{base_prompt}\n\nDocument content:\n\n{extracted_text}", provider="gemini")

        ai_content = clean_ai_html(response)
        now = datetime.now(timezone.utc).isoformat()
        await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {
            "ai_version": ai_content, "ai_processed": True, "rich_content": ai_content, "updated_at": now
        }})
        await log_scheme_action(scheme_id, user["user_id"], user["name"], "ai_digitalized", "File processed by AI")
        return {"ai_content": ai_content}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI processing failed: {str(e)}")
    finally:
        try:
            if tmp_path:
                os.unlink(tmp_path)
        except Exception:
            pass


# ─── UPLOAD SCHEME (creates scheme + attaches file) ───
@router.post("/upload")
async def upload_scheme(data: dict, user: dict = Depends(get_current_user)):
    active_term, academic_year = await get_active_term()
    scheme_id = gen_id()
    now = datetime.now(timezone.utc).isoformat()
    skip_ai = data.get("skip_ai", False)
    doc = {
        "scheme_id": scheme_id,
        "teacher_id": user["user_id"],
        "teacher_name": user.get("name", ""),
        "department": data.get("department", user.get("department", "")),
        "subject_id": data.get("subject_id", ""),
        "subject_name": data.get("subject_name", ""),
        "year_group": data.get("year_group", ""),
        "class_name": data.get("class_name", ""),
        "curriculum": data.get("curriculum", "Cambridge"),
        "term": data.get("term", active_term),
        "academic_year": data.get("academic_year", academic_year),
        "title": data.get("title", f"Uploaded Scheme - {data.get('subject_name', '')}"),
        "weeks": [],
        "rich_content": "<p><em>Original document uploaded — see attached file.</em></p>" if skip_ai else "",
        "original_file_id": data.get("file_id", ""),
        "ai_processed": False,
        "ai_version": None,
        "upload_only": skip_ai,
        "status": "in_progress" if skip_ai else "draft",
        "version": 1,
        "versions": [],
        "reviewer_id": None,
        "reviewer_name": None,
        "approval_notes": "",
        "resubmission_count": 0,
        "created_at": now,
        "updated_at": now,
        "submitted_at": None,
        "approved_at": None,
        "published_at": None,
    }
    await db.schemes.insert_one(doc)
    action_msg = "Original scheme uploaded (no AI)" if skip_ai else "Scheme uploaded from file"
    await log_scheme_action(scheme_id, user["user_id"], user["name"], "uploaded", action_msg)
    doc.pop("_id", None)
    return doc


# ─── DELETE SCHEME ───
@router.delete("/{scheme_id}")
async def delete_scheme(scheme_id: str, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    if scheme["teacher_id"] != user["user_id"] and user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="No permission")
    if scheme["status"] in ("approved", "published"):
        raise HTTPException(status_code=400, detail="Cannot delete approved/published scheme")
    await db.schemes.delete_one({"scheme_id": scheme_id})
    await db.scheme_comments.delete_many({"scheme_id": scheme_id})
    await db.scheme_audit.delete_many({"scheme_id": scheme_id})
    return {"message": "Scheme deleted"}


# ─── WITHDRAW SUBMISSION ───
@router.post("/{scheme_id}/withdraw")
async def withdraw_scheme(scheme_id: str, user: dict = Depends(get_current_user)):
    scheme = await db.schemes.find_one({"scheme_id": scheme_id}, {"_id": 0})
    if not scheme:
        raise HTTPException(status_code=404, detail="Not found")
    if scheme["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your scheme")
    if scheme["status"] not in ("submitted", "resubmitted"):
        raise HTTPException(status_code=400, detail="Can only withdraw submitted schemes")
    now = datetime.now(timezone.utc).isoformat()
    await db.schemes.update_one({"scheme_id": scheme_id}, {"$set": {"status": "in_progress", "updated_at": now}})
    await log_scheme_action(scheme_id, user["user_id"], user["name"], "withdrawn", "Submission withdrawn")
    return {"message": "Submission withdrawn"}
