from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from database import db, get_active_term
from auth import get_current_user
from routes.approval_chain import (
    APPROVAL_CHAIN, build_approval_chain, get_current_step_index,
    is_chain_complete, can_user_approve_step, get_current_step_label,
    can_user_see_lesson, get_user_highest_step,
)
from datetime import datetime, timezone
import jwt as pyjwt
import uuid
import os
import re
import base64

JWT_SECRET = os.environ.get("JWT_SECRET", "default_secret")
JWT_ALGORITHM = "HS256"
MGMT_ROLES = ("SuperAdmin", "Admin", "CEO", "Principal", "HeadTeacher", "Founder", "HeadOfAcademics")

router = APIRouter(prefix="/api/lessons-v2", tags=["lessons-v2"])


def gen_id(prefix="lsn_"):
    return f"{prefix}{uuid.uuid4().hex[:12]}"


async def _get_user_from_token_or_header(request: Request, token: str = None):
    if token:
        try:
            payload = pyjwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
            user_id = payload.get("user_id")
            if user_id:
                user = await db.users.find_one({"user_id": user_id}, {"_id": 0, "password_hash": 0})
                if user:
                    return user
        except Exception:
            pass
    return await get_current_user(request)


STATUS_CONFIG = {
    "draft": "gray",
    "in_progress": "blue",
    "submitted": "indigo",
    "under_review": "amber",
    "returned": "red",
    "resubmitted": "orange",
    "approved": "green",
    "published": "emerald",
    "archived": "slate",
}


async def log_lesson_action(lesson_id, user_id, user_name, action, details=""):
    await db.lesson_audit.insert_one({
        "audit_id": gen_id("laud_"),
        "lesson_id": lesson_id,
        "user_id": user_id,
        "user_name": user_name,
        "action": action,
        "details": details,
        "timestamp": datetime.now(timezone.utc).isoformat(),
    })


# ─── STAFF DASHBOARD ───
@router.get("/my-dashboard")
async def my_dashboard(user: dict = Depends(get_current_user)):
    uid = user["user_id"]
    active_term, _ = await get_active_term()
    total = await db.lessons_v2.count_documents({"teacher_id": uid})
    term_lessons = await db.lessons_v2.count_documents({"teacher_id": uid, "term": active_term})
    drafts = await db.lessons_v2.count_documents({"teacher_id": uid, "status": {"$in": ["draft", "in_progress"]}})
    submitted = await db.lessons_v2.count_documents({"teacher_id": uid, "status": {"$in": ["submitted", "under_review", "resubmitted"]}})
    returned = await db.lessons_v2.count_documents({"teacher_id": uid, "status": "returned"})
    approved = await db.lessons_v2.count_documents({"teacher_id": uid, "status": {"$in": ["approved", "published"]}})
    recent_feedback = await db.lesson_comments.find(
        {
            "lesson_id": {"$in": [s["lesson_id"] async for s in db.lessons_v2.find({"teacher_id": uid}, {"lesson_id": 1})]},
            "is_management": True,
        },
        {"_id": 0},
    ).sort("timestamp", -1).to_list(5)
    return {
        "total": total, "term_lessons": term_lessons, "drafts": drafts,
        "submitted": submitted, "returned": returned, "approved": approved,
        "active_term": active_term, "recent_feedback": recent_feedback,
    }


# ─── MANAGEMENT DASHBOARD ───
@router.get("/management-dashboard")
async def management_dashboard(user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management access only")
    active_term, _ = await get_active_term()
    # Fetch all non-draft lessons and filter by chain visibility
    all_lessons = await db.lessons_v2.find({"status": {"$ne": "draft"}}, {"_id": 0, "status": 1, "teacher_id": 1, "approval_chain": 1, "term": 1}).to_list(1000)
    visible = [l for l in all_lessons if l.get("teacher_id") == user["user_id"] or can_user_see_lesson(user, l.get("approval_chain", []))]
    total = len(visible)
    pending = sum(1 for l in visible if l["status"] in ("submitted", "resubmitted"))
    under_review = sum(1 for l in visible if l["status"] == "under_review")
    approved_count = sum(1 for l in visible if l["status"] in ("approved", "published"))
    returned_count = sum(1 for l in visible if l["status"] == "returned")
    # Staff who haven't submitted
    all_teachers = await db.users.find(
        {"role": {"$nin": ["Student", "Parent"]}, "roles": {"$in": ["Teacher"]}},
        {"_id": 0, "user_id": 1, "name": 1, "department": 1, "teaching_subjects": 1},
    ).to_list(500)
    submitted_teachers = set()
    async for s in db.lessons_v2.find({"term": active_term, "status": {"$nin": ["draft"]}}, {"teacher_id": 1}):
        submitted_teachers.add(s["teacher_id"])
    overdue_staff = [t for t in all_teachers if t["user_id"] not in submitted_teachers]
    return {
        "total": total, "pending": pending, "under_review": under_review,
        "approved": approved_count, "returned": returned_count,
        "overdue_staff": overdue_staff, "active_term": active_term,
    }


# ─── LIST LESSONS ───
@router.get("/")
async def list_lessons(
    user: dict = Depends(get_current_user),
    status: str = Query(None),
    department: str = Query(None),
    subject: str = Query(None),
    teacher_id: str = Query(None),
    term: str = Query(None),
    year_group: str = Query(None),
):
    is_mgmt = user.get("role") in MGMT_ROLES
    query = {}
    if not is_mgmt:
        query["teacher_id"] = user["user_id"]
    else:
        if status == "all_submitted":
            query["status"] = {"$ne": "draft"}
        elif status:
            query["status"] = status
        if teacher_id:
            query["teacher_id"] = teacher_id
    if department:
        query["department"] = department
    if subject:
        query["subject_id"] = subject
    if term:
        query["term"] = term
    if year_group:
        query["year_group"] = year_group
    if not is_mgmt and status:
        query["status"] = status
    lessons = await db.lessons_v2.find(query, {"_id": 0}).sort("updated_at", -1).to_list(500)
    # Filter by approval chain visibility for management
    if is_mgmt:
        lessons = [l for l in lessons if l.get("teacher_id") == user["user_id"] or can_user_see_lesson(user, l.get("approval_chain", []))]
    return lessons


# ─── CREATE LESSON (from scratch) ───
@router.post("/")
async def create_lesson(data: dict, user: dict = Depends(get_current_user)):
    active_term, academic_year = await get_active_term()
    lesson_id = gen_id()
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "lesson_id": lesson_id,
        "teacher_id": user["user_id"],
        "teacher_name": user.get("name", ""),
        "department": data.get("department", user.get("department", "")),
        "subject_id": data.get("subject_id", ""),
        "subject_name": data.get("subject_name", ""),
        "year_group": data.get("year_group", ""),
        "class_ids": data.get("class_ids", []),
        "curriculum": data.get("curriculum", "Cambridge"),
        "term": data.get("term", active_term),
        "academic_year": data.get("academic_year", academic_year),
        "title": data.get("title", "Untitled Lesson"),
        "description": data.get("description", ""),
        "objectives": data.get("objectives", ""),
        "rich_content": data.get("rich_content", ""),
        "content_text": data.get("content_text", ""),
        "duration_minutes": data.get("duration_minutes", 45),
        "lesson_type": data.get("lesson_type", "mixed"),
        "file_ids": data.get("file_ids", []),
        "content_url": data.get("content_url", ""),
        "extra_links": data.get("extra_links", []),
        "tags": data.get("tags", []),
        "original_file_id": None,
        "ai_processed": False,
        "ai_version": None,
        "status": "draft",
        "version": 1,
        "versions": [],
        "reviewer_id": None,
        "reviewer_name": None,
        "approval_notes": "",
        "resubmission_count": 0,
        "created_at": now,
        "updated_at": now,
        "submitted_at": None,
        "approved_at": None,
        "published_at": None,
    }
    await db.lessons_v2.insert_one(doc)
    template_id = data.get("template_id")
    if template_id:
        await db.lesson_templates.update_one({"template_id": template_id}, {"$inc": {"usage_count": 1}})
    await log_lesson_action(lesson_id, user["user_id"], user["name"], "created",
                            f"Lesson created{' from template' if template_id else ' from scratch'}")
    doc.pop("_id", None)
    return doc


# ─── UPLOAD LESSON (creates lesson + attaches file) ───
@router.post("/upload")
async def upload_lesson(data: dict, user: dict = Depends(get_current_user)):
    active_term, academic_year = await get_active_term()
    lesson_id = gen_id()
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "lesson_id": lesson_id,
        "teacher_id": user["user_id"],
        "teacher_name": user.get("name", ""),
        "department": data.get("department", user.get("department", "")),
        "subject_id": data.get("subject_id", ""),
        "subject_name": data.get("subject_name", ""),
        "year_group": data.get("year_group", ""),
        "class_ids": data.get("class_ids", []),
        "curriculum": data.get("curriculum", "Cambridge"),
        "term": data.get("term", active_term),
        "academic_year": data.get("academic_year", academic_year),
        "title": data.get("title", "Uploaded Lesson"),
        "description": data.get("description", ""),
        "objectives": data.get("objectives", ""),
        "rich_content": "",
        "content_text": "",
        "duration_minutes": data.get("duration_minutes", 45),
        "lesson_type": data.get("lesson_type", "mixed"),
        "file_ids": data.get("file_ids", []),
        "content_url": "",
        "extra_links": [],
        "tags": [],
        "original_file_id": data.get("file_id", ""),
        "ai_processed": False,
        "ai_version": None,
        "status": "draft",
        "version": 1,
        "versions": [],
        "reviewer_id": None,
        "reviewer_name": None,
        "approval_notes": "",
        "resubmission_count": 0,
        "created_at": now,
        "updated_at": now,
        "submitted_at": None,
        "approved_at": None,
        "published_at": None,
    }
    await db.lessons_v2.insert_one(doc)
    await log_lesson_action(lesson_id, user["user_id"], user["name"], "uploaded", "Lesson uploaded from file")
    doc.pop("_id", None)
    return doc


# ─── MY PENDING REVIEWS (must be before /{lesson_id}) ───
@router.get("/my-pending-reviews")
async def my_pending_reviews(user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    all_pending = await db.lessons_v2.find(
        {"status": {"$in": ["submitted", "under_review", "resubmitted"]}, "approval_chain": {"$exists": True}},
        {"_id": 0},
    ).sort("submitted_at", -1).to_list(200)
    my_items = []
    for item in all_pending:
        chain = item.get("approval_chain", [])
        idx = get_current_step_index(chain)
        if idx >= 0 and can_user_approve_step(user, chain[idx], chain):
            my_items.append(item)
    return my_items


# ─── GET SINGLE LESSON ───
@router.get("/{lesson_id}")
async def get_lesson(lesson_id: str, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Lesson not found")
    comments = await db.lesson_comments.find({"lesson_id": lesson_id}, {"_id": 0}).sort("timestamp", 1).to_list(200)
    audit = await db.lesson_audit.find({"lesson_id": lesson_id}, {"_id": 0}).sort("timestamp", -1).to_list(100)
    return {**lesson, "comments": comments, "audit_trail": audit}


# ─── UPDATE LESSON ───
@router.put("/{lesson_id}")
async def update_lesson(lesson_id: str, data: dict, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    if lesson["teacher_id"] != user["user_id"] and user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Not your lesson")
    if lesson["status"] in ("approved", "published", "archived"):
        raise HTTPException(status_code=400, detail="Cannot edit approved/published lesson")
    safe = {
        "title", "description", "objectives", "department", "subject_id", "subject_name",
        "year_group", "class_ids", "curriculum", "term", "academic_year",
        "rich_content", "content_text", "duration_minutes", "lesson_type",
        "file_ids", "content_url", "extra_links", "tags", "original_file_id",
    }
    update = {k: v for k, v in data.items() if k in safe}
    update["updated_at"] = datetime.now(timezone.utc).isoformat()
    if lesson["status"] == "draft":
        update["status"] = "in_progress"
    await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": update})
    await log_lesson_action(lesson_id, user["user_id"], user["name"], "edited", "Content updated")
    updated = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    return updated


# ─── SUBMIT TO MANAGEMENT ───
@router.post("/{lesson_id}/submit")
async def submit_lesson(lesson_id: str, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    if lesson["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your lesson")
    if lesson["status"] not in ("draft", "in_progress", "returned"):
        raise HTTPException(status_code=400, detail=f"Cannot submit from '{lesson['status']}' status")
    now = datetime.now(timezone.utc).isoformat()
    new_status = "resubmitted" if lesson["status"] == "returned" else "submitted"
    chain = build_approval_chain(user)
    # If all steps auto-approved (e.g. Founder submits), mark as approved
    if is_chain_complete(chain):
        new_status = "approved"
    current_idx = get_current_step_index(chain)
    update = {
        "status": new_status, "submitted_at": now, "updated_at": now,
        "approval_chain": chain,
        "current_approval_step": chain[current_idx]["step"] if current_idx >= 0 else 0,
        "current_approver_label": get_current_step_label(chain),
    }
    if new_status == "approved":
        update["approved_at"] = now
    if lesson.get("status") == "returned":
        update["resubmission_count"] = lesson.get("resubmission_count", 0) + 1
    version_snap = {
        "version": lesson.get("version", 1),
        "rich_content": lesson.get("rich_content", ""),
        "saved_at": now,
    }
    await db.lessons_v2.update_one(
        {"lesson_id": lesson_id},
        {"$set": update, "$push": {"versions": version_snap}, "$inc": {"version": 1}},
    )
    await log_lesson_action(lesson_id, user["user_id"], user["name"], new_status,
                            f"Lesson {new_status} — {get_current_step_label(chain)}")
    return {"message": f"Lesson {new_status} successfully", "status": new_status, "current_approver": get_current_step_label(chain)}


# ─── WITHDRAW SUBMISSION ───
@router.post("/{lesson_id}/withdraw")
async def withdraw_lesson(lesson_id: str, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    if lesson["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your lesson")
    if lesson["status"] not in ("submitted", "resubmitted"):
        raise HTTPException(status_code=400, detail="Can only withdraw submitted lessons")
    now = datetime.now(timezone.utc).isoformat()
    await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {"status": "in_progress", "updated_at": now}})
    await log_lesson_action(lesson_id, user["user_id"], user["name"], "withdrawn", "Submission withdrawn")
    return {"message": "Submission withdrawn"}


# ─── MANAGEMENT: REVIEW ACTION (Hierarchical) ───
@router.post("/{lesson_id}/review")
async def review_lesson(lesson_id: str, data: dict, user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    action = data.get("action")
    comment_text = data.get("comment", "")
    now = datetime.now(timezone.utc).isoformat()
    chain = lesson.get("approval_chain", [])

    if action in ("approve", "return") and chain:
        # Find the current pending step
        current_idx = get_current_step_index(chain)
        if current_idx == -1:
            raise HTTPException(status_code=400, detail="All steps already completed")
        current_step = chain[current_idx]
        if not can_user_approve_step(user, current_step, chain):
            raise HTTPException(status_code=403,
                                detail=f"This item is currently with {current_step['label']}. You cannot act on it at this stage.")

        if action == "approve":
            chain[current_idx]["status"] = "approved"
            chain[current_idx]["user_id"] = user["user_id"]
            chain[current_idx]["user_name"] = user.get("name", "")
            chain[current_idx]["acted_at"] = now
            chain[current_idx]["comment"] = comment_text

            if is_chain_complete(chain):
                await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
                    "status": "approved", "approved_at": now, "approval_notes": comment_text,
                    "approval_chain": chain, "current_approval_step": 0,
                    "current_approver_label": "Fully Approved",
                    "reviewer_id": user["user_id"], "reviewer_name": user.get("name", ""),
                    "updated_at": now,
                }})
                await log_lesson_action(lesson_id, user["user_id"], user["name"], "approved",
                                        f"Final approval by {current_step['label']}. {comment_text}")
            else:
                next_idx = get_current_step_index(chain)
                next_label = get_current_step_label(chain)
                await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
                    "status": "under_review", "approval_chain": chain,
                    "current_approval_step": chain[next_idx]["step"],
                    "current_approver_label": next_label,
                    "reviewer_id": user["user_id"], "reviewer_name": user.get("name", ""),
                    "updated_at": now,
                }})
                await log_lesson_action(lesson_id, user["user_id"], user["name"], "step_approved",
                                        f"Approved by {current_step['label']} → Now {next_label}. {comment_text}")

        elif action == "return":
            chain[current_idx]["status"] = "returned"
            chain[current_idx]["user_id"] = user["user_id"]
            chain[current_idx]["user_name"] = user.get("name", "")
            chain[current_idx]["acted_at"] = now
            chain[current_idx]["comment"] = comment_text
            await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
                "status": "returned", "approval_chain": chain,
                "current_approval_step": current_step["step"],
                "current_approver_label": f"Returned by {current_step['label']}",
                "reviewer_id": user["user_id"], "reviewer_name": user.get("name", ""),
                "updated_at": now,
            }})
            await log_lesson_action(lesson_id, user["user_id"], user["name"], "returned",
                                    f"Returned by {current_step['label']}. {comment_text}")

    elif action == "start_review":
        await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
            "status": "under_review", "reviewer_id": user["user_id"],
            "reviewer_name": user.get("name", ""), "updated_at": now,
        }})
        await log_lesson_action(lesson_id, user["user_id"], user["name"], "under_review", "Review started")

    elif action == "publish":
        await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
            "status": "published", "published_at": now, "updated_at": now,
        }})
        await log_lesson_action(lesson_id, user["user_id"], user["name"], "published", "Lesson published")

    elif action == "archive":
        await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
            "status": "archived", "updated_at": now,
        }})
        await log_lesson_action(lesson_id, user["user_id"], user["name"], "archived", "Lesson archived")

    if comment_text:
        await db.lesson_comments.insert_one({
            "comment_id": gen_id("lcmt_"), "lesson_id": lesson_id,
            "user_id": user["user_id"], "user_name": user.get("name", ""),
            "text": comment_text, "is_management": True, "action": action,
            "timestamp": now,
        })
    return {"message": f"Action '{action}' completed", "status": action}


# ─── MY PENDING REVIEWS ───


# ─── ADD COMMENT ───
@router.post("/{lesson_id}/comments")
async def add_comment(lesson_id: str, data: dict, user: dict = Depends(get_current_user)):
    is_mgmt = user.get("role") in MGMT_ROLES
    now = datetime.now(timezone.utc).isoformat()
    await db.lesson_comments.insert_one({
        "comment_id": gen_id("lcmt_"), "lesson_id": lesson_id,
        "user_id": user["user_id"], "user_name": user.get("name", ""),
        "text": data.get("text", ""), "is_management": is_mgmt,
        "section": data.get("section", ""), "timestamp": now,
    })
    await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {"updated_at": now}})
    return {"message": "Comment added"}


# ─── COMPLIANCE TRACKER ───
@router.get("/management/compliance")
async def compliance_tracker(user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    active_term, _ = await get_active_term()
    teachers = await db.users.find(
        {"roles": {"$in": ["Teacher"]}},
        {"_id": 0, "user_id": 1, "name": 1, "department": 1, "teaching_subjects": 1},
    ).to_list(500)
    lessons = await db.lessons_v2.find(
        {"term": active_term}, {"_id": 0, "teacher_id": 1, "subject_id": 1, "status": 1}
    ).to_list(1000)
    lesson_map = {}
    for rec in lessons:
        key = rec["teacher_id"]
        if key not in lesson_map:
            lesson_map[key] = []
        lesson_map[key].append(rec)
    result = []
    for t in teachers:
        t_lessons = lesson_map.get(t["user_id"], [])
        subjects = t.get("teaching_subjects", [])
        required = max(len(subjects), 1)
        submitted_count = sum(1 for rec in t_lessons if rec["status"] not in ("draft",))
        approved_count = sum(1 for rec in t_lessons if rec["status"] in ("approved", "published"))
        result.append({
            "user_id": t["user_id"], "name": t["name"],
            "department": t.get("department", ""), "subjects": subjects,
            "required": required, "submitted": submitted_count,
            "approved": approved_count, "pending": submitted_count - approved_count,
            "missing": max(required - submitted_count, 0),
            "rate": round(submitted_count / required * 100) if required > 0 else 0,
        })
    result.sort(key=lambda x: x["rate"])
    return result


# ─── LESSON TEMPLATES ───
@router.get("/templates/list")
async def list_templates(
    subject: str = Query(None),
    year_group: str = Query(None),
    curriculum: str = Query(None),
    user: dict = Depends(get_current_user),
):
    query = {"is_active": True}
    if subject:
        query["subject_name"] = {"$regex": subject, "$options": "i"}
    if year_group:
        query["year_group"] = year_group
    if curriculum:
        query["curriculum"] = curriculum
    templates = await db.lesson_templates.find(query, {"_id": 0}).sort("created_at", -1).to_list(200)
    return templates


@router.post("/templates/create")
async def create_template(data: dict, user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    now = datetime.now(timezone.utc).isoformat()
    template_id = gen_id("ltpl_")
    doc = {
        "template_id": template_id,
        "name": data.get("name", "Untitled Template"),
        "description": data.get("description", ""),
        "subject_name": data.get("subject_name", ""),
        "year_group": data.get("year_group", ""),
        "curriculum": data.get("curriculum", "Cambridge"),
        "duration_minutes": data.get("duration_minutes", 45),
        "objectives": data.get("objectives", ""),
        "rich_content": data.get("rich_content", ""),
        "tags": data.get("tags", []),
        "is_active": True,
        "created_by": user["user_id"],
        "created_by_name": user.get("name", ""),
        "created_at": now,
        "updated_at": now,
        "usage_count": 0,
    }
    await db.lesson_templates.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.put("/templates/{template_id}")
async def update_template(template_id: str, data: dict, user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    template = await db.lesson_templates.find_one({"template_id": template_id}, {"_id": 0})
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    safe = {"name", "description", "subject_name", "year_group", "curriculum",
            "duration_minutes", "objectives", "rich_content", "tags", "is_active"}
    update = {k: v for k, v in data.items() if k in safe}
    update["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.lesson_templates.update_one({"template_id": template_id}, {"$set": update})
    updated = await db.lesson_templates.find_one({"template_id": template_id}, {"_id": 0})
    return updated


@router.delete("/templates/{template_id}")
async def delete_template(template_id: str, user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    result = await db.lesson_templates.delete_one({"template_id": template_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Template not found")
    return {"message": "Template deleted"}


# ─── AI GENERATE TEMPLATE ───
@router.post("/templates/ai-generate")
async def ai_generate_template(data: dict, user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    subject = data.get("subject_name", "General")
    year_group = data.get("year_group", "")
    topic = data.get("topic", "")
    curriculum = data.get("curriculum", "Cambridge")
    duration = data.get("duration_minutes", 45)
    lesson_type = data.get("lesson_type", "mixed")
    custom_instructions = data.get("custom_instructions", "")

    type_label = {"mixed": "mixed theory and practice", "theory": "theory-based lecture",
                  "practical": "practical/hands-on", "assessment": "assessment-focused",
                  "revision": "revision/review"}.get(lesson_type, "mixed")

    try:
        system = f"""You are a senior curriculum designer for The RightWay Centre of Learning, a prestigious {curriculum} curriculum school.
Generate a COMPLETE, professional lesson plan template that teachers can reuse.

REQUIREMENTS:
- Subject: {subject}
{"- Year Group: " + year_group if year_group else "- Year Group: Generic (suitable for any year)"}
{"- Topic Focus: " + topic if topic else ""}
- Duration: {duration} minutes
- Lesson Type: {type_label}
{f"- Additional: {custom_instructions}" if custom_instructions else ""}

OUTPUT FORMAT — Return a JSON object with these fields:
{{
  "name": "<template name, e.g. 'RightWay {subject} Practical Template'>",
  "description": "<1-2 sentence description of the template>",
  "objectives": "<3-4 SMART learning objectives using Bloom's taxonomy verbs>",
  "rich_content": "<COMPLETE HTML lesson plan structure>"
}}

The rich_content HTML MUST include these sections:
1. <h2>Lesson Plan: [Topic]</h2>
2. <h3>1. Starter Activity / Hook (X min)</h3> — with example activities
3. <h3>2. Main Teaching Phase (X min)</h3> — with key teaching points
4. <h3>3. Student Activities (X min)</h3> — with differentiated tasks
5. <h3>4. Plenary / Wrap-up (X min)</h3> — with assessment strategies
6. <h3>5. Assessment & Success Criteria</h3> — with formative checks
7. <h3>6. Differentiation</h3> — EAL, SEN, Gifted strategies
8. <h3>7. Resources Needed</h3> — materials list
9. <h3>8. Homework / Extension</h3>

Use proper HTML tables where appropriate. Include placeholder text like [TEACHER: Add specific content here].
Timing in each section header must add up to {duration} minutes.
No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags. Return ONLY the JSON object, no markdown fences."""

        from services.ai_service import llm_chat
        response = await llm_chat(system, f"Generate a professional {subject} lesson plan template{' for ' + year_group if year_group else ''}{' on ' + topic if topic else ''} ({duration} min, {type_label}).", provider="gemini"
        )

        import json
        text = response.strip()
        text = re.sub(r'^```(?:json)?\s*\n?', '', text)
        text = re.sub(r'\n?```\s*$', '', text)
        parsed = json.loads(text)

        now = datetime.now(timezone.utc).isoformat()
        template_id = gen_id("ltpl_")
        doc = {
            "template_id": template_id,
            "name": parsed.get("name", f"AI {subject} Template"),
            "description": parsed.get("description", f"AI-generated {subject} lesson plan template"),
            "subject_name": subject,
            "year_group": year_group,
            "curriculum": curriculum,
            "duration_minutes": duration,
            "objectives": parsed.get("objectives", ""),
            "rich_content": _clean_ai_html(parsed.get("rich_content", "")),
            "tags": ["ai-generated"],
            "is_active": True,
            "created_by": user["user_id"],
            "created_by_name": user.get("name", ""),
            "created_at": now,
            "updated_at": now,
            "usage_count": 0,
        }
        await db.lesson_templates.insert_one(doc)
        doc.pop("_id", None)
        return doc
    except json.JSONDecodeError:
        raise HTTPException(status_code=500, detail="AI returned invalid format. Try again.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI generation failed: {str(e)}")


# ─── AI IMPROVE TEMPLATE ───
@router.post("/templates/{template_id}/ai-improve")
async def ai_improve_template(template_id: str, data: dict, user: dict = Depends(get_current_user)):
    if user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="Management only")
    template = await db.lesson_templates.find_one({"template_id": template_id}, {"_id": 0})
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    content = template.get("rich_content", "")
    if not content or len(content) < 20:
        raise HTTPException(status_code=400, detail="No content to improve")

    prompt_type = data.get("prompt_type", "improve")
    custom_prompt = data.get("custom_prompt", "")

    ctx = (
        f"Subject: {template.get('subject_name', 'General')}\n"
        f"Year Group: {template.get('year_group', 'All')}\n"
        f"Duration: {template.get('duration_minutes', 45)} minutes\n"
        f"Objectives: {template.get('objectives', '')}\n"
    )

    TABLE_RULE = "\n\nPreserve all HTML tables exactly. Return COMPLETE HTML. No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags."
    prompts = {
        "improve": "You are an expert curriculum designer for TCIS school. Improve this lesson plan template: add more detail, strengthen placeholder instructions, ensure professional structure." + TABLE_RULE,
        "standardise": "Rewrite this template in strict TCIS standard format with sections: Starter/Hook, Main Teaching, Student Activities, Plenary, Assessment, Differentiation, Resources, Homework." + TABLE_RULE,
        "expand": "Expand all placeholder sections with detailed example activities, questions, and teaching strategies that teachers can adapt." + TABLE_RULE,
        "differentiation": "Add comprehensive differentiation strategies throughout: EAL learners, SEN students, gifted & talented, visual/auditory/kinesthetic learners." + TABLE_RULE,
        "assessment": "Add or improve assessment elements: formative checks, exit tickets, success criteria ladders, questioning frameworks, peer/self-assessment tasks." + TABLE_RULE,
        "engagement": "Add creative engagement strategies: interactive activities, think-pair-share, mini-whiteboards, student-led segments, games, real-world connections." + TABLE_RULE,
        "custom": (custom_prompt or "Improve this template.") + TABLE_RULE,
    }
    system_prompt = prompts.get(prompt_type, prompts["improve"])

    try:
        from services.ai_service import llm_chat
        response = await llm_chat(system_prompt, f"{ctx}\n\nTemplate content:\n\n{content}", provider="gemini")
        ai_content = _clean_ai_html(response)
        now = datetime.now(timezone.utc).isoformat()
        await db.lesson_templates.update_one({"template_id": template_id}, {"$set": {
            "rich_content": ai_content, "updated_at": now,
        }})
        return {"ai_content": ai_content, "prompt_type": prompt_type}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI improvement failed: {str(e)}")


# ─── DELETE LESSON ───
@router.delete("/{lesson_id}")
async def delete_lesson(lesson_id: str, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    if lesson["teacher_id"] != user["user_id"] and user.get("role") not in MGMT_ROLES:
        raise HTTPException(status_code=403, detail="No permission")
    if lesson["status"] in ("approved", "published"):
        raise HTTPException(status_code=400, detail="Cannot delete approved/published lesson")
    await db.lessons_v2.delete_one({"lesson_id": lesson_id})
    await db.lesson_comments.delete_many({"lesson_id": lesson_id})
    await db.lesson_audit.delete_many({"lesson_id": lesson_id})
    return {"message": "Lesson deleted"}


# ─── AI GENERATE LESSON (for teachers) ───
@router.post("/ai-generate-lesson")
async def ai_generate_lesson(data: dict, user: dict = Depends(get_current_user)):
    """Generate a full lesson plan from scratch using AI. Available to all staff."""
    subject = data.get("subject", "General")
    year_group = data.get("year_group", "")
    topic = data.get("topic", "")
    curriculum = data.get("curriculum", "Cambridge")
    duration = data.get("duration_minutes", 45)
    additional_instructions = data.get("additional_instructions", "")

    if not topic.strip():
        raise HTTPException(status_code=400, detail="Topic is required for AI generation")

    try:
        from services.ai_service import llm_chat
        system = f"""You are a senior curriculum designer for The RightWay Centre of Learning, a prestigious {curriculum} curriculum school.
Generate a COMPLETE, professional lesson plan.

REQUIREMENTS:
- Subject: {subject}
- Year Group: {year_group or 'Not specified'}
- Topic: {topic}
- Duration: {duration} minutes
{f"- Additional: {additional_instructions}" if additional_instructions else ""}

OUTPUT FORMAT — Return a JSON object:
{{
  "title": "<lesson title, e.g. 'Photosynthesis: The Light Reactions'>",
  "description": "<1-2 sentence description>",
  "objectives": "<3-4 SMART learning objectives using Bloom's taxonomy>",
  "rich_content": "<COMPLETE HTML lesson plan>"
}}

The rich_content HTML MUST include:
1. <h2>Lesson Plan: [Topic]</h2>
2. <h3>1. Starter Activity / Hook (X min)</h3> — engaging opener
3. <h3>2. Main Teaching Phase (X min)</h3> — key teaching points, explanations
4. <h3>3. Student Activities (X min)</h3> — differentiated tasks for varied abilities
5. <h3>4. Plenary / Wrap-up (X min)</h3> — consolidation, exit ticket
6. <h3>5. Assessment & Success Criteria</h3> — formative checks
7. <h3>6. Differentiation</h3> — EAL, SEN, Gifted & Talented
8. <h3>7. Resources Needed</h3> — materials list
9. <h3>8. Homework / Extension</h3>

Use proper HTML tables where appropriate. Timing must add up to {duration} minutes.
No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags. Return ONLY the JSON object, no markdown fences."""

        response = await llm_chat(system, f"Generate a {subject} lesson plan for {year_group} on '{topic}' ({duration} minutes).", provider="gemini")

        import json
        text = response.strip()
        text = re.sub(r'^```(?:json)?\s*\n?', '', text)
        text = re.sub(r'\n?```\s*$', '', text)
        parsed = json.loads(text)

        return {
            "title": parsed.get("title", f"{topic} - {subject}"),
            "description": parsed.get("description", ""),
            "objectives": parsed.get("objectives", ""),
            "rich_content": _clean_ai_html(parsed.get("rich_content", "")),
        }
    except json.JSONDecodeError:
        raise HTTPException(status_code=500, detail="AI returned invalid format. Please try again.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI generation failed: {str(e)}")


# ─── AI IMPROVE LESSON ───
@router.post("/{lesson_id}/ai-improve")
async def ai_improve_lesson(lesson_id: str, data: dict, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    if lesson["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your lesson")
    prompt_type = data.get("prompt_type", "improve")
    custom_prompt = data.get("custom_prompt", "")
    content = lesson.get("rich_content", "") or lesson.get("content_text", "")
    full_content = (
        f"Subject: {lesson.get('subject_name', '')}\n"
        f"Year Group: {lesson.get('year_group', '')}\n"
        f"Term: {lesson.get('term', '')}\n"
        f"Objectives: {lesson.get('objectives', '')}\n"
        f"Duration: {lesson.get('duration_minutes', 45)} minutes\n\n"
        f"{content}"
    )
    TABLE_RULE = "\n\nIf the content contains an HTML table, preserve the exact <table>, <thead>, <tr>, <th>, <tbody>, <td> structure. Return COMPLETE table. No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags."
    prompts = {
        "improve": "You are an expert lesson planning assistant for TCIS (The RightWay Centre of Learning). Improve this lesson plan professionally. Strengthen learning objectives, improve clarity, ensure alignment between objectives, activities, and assessments. Format as a clear, structured HTML document." + TABLE_RULE,
        "standardise": "Rewrite this lesson plan in TCIS standard format. Structure clearly with: Learning Objectives, Starter Activity, Main Teaching, Student Activities, Plenary, Assessment, Homework, Resources Needed." + TABLE_RULE,
        "expand": "Expand weak or incomplete sections of this lesson plan. Add more detail to activities, differentiation strategies, and assessment methods." + TABLE_RULE,
        "objectives": "Rewrite all learning objectives using Bloom's taxonomy measurable verbs (analyze, evaluate, create, apply, etc.). Make each objective SMART." + TABLE_RULE,
        "assessment": "Add or improve assessment strategies. Suggest formative checks, exit tickets, and differentiated questioning." + TABLE_RULE,
        "differentiation": "Add differentiation strategies throughout this lesson for EAL learners, SEN students, and gifted/talented students." + TABLE_RULE,
        "digitalize": "Convert this content into a clean, structured HTML lesson plan with sections: Learning Objectives, Starter/Hook, Main Teaching Phase, Student Activities, Plenary/Wrap-up, Assessment, Homework, Resources." + TABLE_RULE,
        "custom": (custom_prompt or "Improve this lesson plan.") + TABLE_RULE,
    }
    system_prompt = prompts.get(prompt_type, prompts["improve"])
    try:
        from services.ai_service import llm_chat
        response = await llm_chat(system_prompt, f"Here is the lesson plan content:\n\n{full_content}", provider="gemini")
        ai_content = _clean_ai_html(response)
        now = datetime.now(timezone.utc).isoformat()
        await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
            "ai_version": ai_content, "ai_processed": True,
            "rich_content": ai_content, "updated_at": now,
        }})
        await log_lesson_action(lesson_id, user["user_id"], user["name"], "ai_improved", f"AI {prompt_type}")
        return {"ai_content": ai_content, "prompt_type": prompt_type}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI processing failed: {str(e)}")


# ─── AI ANALYZE LESSON ───
@router.post("/{lesson_id}/ai-analyze")
async def ai_analyze_lesson(lesson_id: str, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    if lesson["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your lesson")
    content = lesson.get("rich_content", "") or lesson.get("content_text", "")
    if not content or len(content) < 20:
        return {"issues": []}
    try:
        from services.ai_service import llm_chat
        system = """You are a strict academic quality reviewer for The RightWay Centre of Learning.
Analyze this lesson plan and return ONLY a valid JSON array of issues. Each issue object must have:
- "id": a short unique slug (e.g. "weak-objectives")
- "category": one of "Objectives", "Activities", "Assessment", "Differentiation", "Resources", "Timing", "Engagement", "Alignment"
- "severity": one of "high", "medium", "low"
- "title": a short 5-10 word summary
- "description": 1-2 sentences explaining the issue and fix
Return ONLY the JSON array, no markdown, no code fences."""
        response = await llm_chat(system, f"Subject: {lesson.get('subject_name','')}\nYear: {lesson.get('year_group','')}\nObjectives: {lesson.get('objectives','')}\n\nLesson content:\n{content}", provider="gemini")
        import json
        text = response.strip()
        text = re.sub(r'^```(?:json)?\s*\n?', '', text)
        text = re.sub(r'\n?```\s*$', '', text)
        issues = json.loads(text)
        if not isinstance(issues, list):
            issues = []
        await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {"ai_issues": issues}})
        return {"issues": issues}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Analysis failed: {str(e)}")


# ─── AI FIX SPECIFIC ISSUE ───
@router.post("/{lesson_id}/ai-fix-issue")
async def ai_fix_issue(lesson_id: str, data: dict, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    if lesson["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your lesson")
    issue_title = data.get("issue_title", "")
    issue_description = data.get("issue_description", "")
    issue_id = data.get("issue_id", "")
    content = lesson.get("rich_content", "")
    if not content:
        raise HTTPException(status_code=400, detail="No content to fix")
    try:
        from services.ai_service import llm_chat
        system = f"""You are an expert lesson planning assistant for The RightWay Centre of Learning.
Fix ONLY this specific issue: "{issue_title}" - {issue_description}
Keep the same HTML structure. Only modify content related to the issue.
Return the COMPLETE HTML. No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags."""
        response = await llm_chat(system, f"Lesson plan:\n\n{content}", provider="gemini")
        ai_content = _clean_ai_html(response)
        now = datetime.now(timezone.utc).isoformat()
        existing_issues = lesson.get("ai_issues", [])
        remaining_issues = [i for i in existing_issues if i.get("id") != issue_id]
        await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
            "rich_content": ai_content, "ai_version": ai_content,
            "ai_processed": True, "ai_issues": remaining_issues, "updated_at": now,
        }})
        await log_lesson_action(lesson_id, user["user_id"], user["name"], "ai_fix_issue", f"Fixed: {issue_title}")
        return {"ai_content": ai_content, "remaining_issues": remaining_issues}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Fix failed: {str(e)}")


# ─── AI FIX ALL ISSUES ───
@router.post("/{lesson_id}/ai-fix-all")
async def ai_fix_all_issues(lesson_id: str, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    if lesson["teacher_id"] != user["user_id"]:
        raise HTTPException(status_code=403, detail="Not your lesson")
    issues = lesson.get("ai_issues", [])
    content = lesson.get("rich_content", "")
    if not content or not issues:
        raise HTTPException(status_code=400, detail="No content or no issues to fix")
    issues_text = "\n".join([f"- [{i.get('severity','medium').upper()}] {i.get('title','')}: {i.get('description','')}" for i in issues])
    try:
        from services.ai_service import llm_chat
        system = f"""Fix ALL of the following issues in the lesson plan:
{issues_text}
Keep the same HTML structure. Return COMPLETE HTML. No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags."""
        response = await llm_chat(system, f"Lesson plan:\n\n{content}", provider="gemini")
        ai_content = _clean_ai_html(response)
        now = datetime.now(timezone.utc).isoformat()
        await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
            "rich_content": ai_content, "ai_version": ai_content,
            "ai_processed": True, "ai_issues": [], "updated_at": now,
        }})
        await log_lesson_action(lesson_id, user["user_id"], user["name"], "ai_fix_all", f"Fixed {len(issues)} issues")
        return {"ai_content": ai_content, "fixed_count": len(issues)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Fix all failed: {str(e)}")


# ─── AI EVALUATE LESSON (Management) ───
@router.post("/{lesson_id}/ai-evaluate")
async def ai_evaluate_lesson(lesson_id: str, user: dict = Depends(get_current_user)):
    if not any(r in user.get("roles", [user.get("role", "")]) for r in MGMT_ROLES):
        raise HTTPException(status_code=403, detail="Management only")
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    content = lesson.get("rich_content", "") or lesson.get("content_text", "")
    if not content or len(content) < 20:
        raise HTTPException(status_code=400, detail="Not enough content to evaluate")
    try:
        from services.ai_service import llm_chat
        system = """You are a senior academic quality evaluator for The RightWay Centre of Learning.
Grade this lesson plan and return ONLY a valid JSON object with:
{
  "overall_score": <0-100>,
  "grade": "<A+/A/B+/B/C+/C/D/F>",
  "summary": "<2-3 sentence assessment>",
  "categories": [
    {"name": "Learning Objectives", "score": <0-100>, "feedback": "<feedback>"},
    {"name": "Teaching Activities", "score": <0-100>, "feedback": "<feedback>"},
    {"name": "Student Engagement", "score": <0-100>, "feedback": "<feedback>"},
    {"name": "Assessment Strategy", "score": <0-100>, "feedback": "<feedback>"},
    {"name": "Differentiation", "score": <0-100>, "feedback": "<feedback>"},
    {"name": "Resource Quality", "score": <0-100>, "feedback": "<feedback>"},
    {"name": "Timing & Pacing", "score": <0-100>, "feedback": "<feedback>"},
    {"name": "Completeness", "score": <0-100>, "feedback": "<feedback>"}
  ],
  "errors": [{"issue": "<desc>", "severity": "<high/medium/low>", "suggestion": "<fix>"}],
  "strengths": ["<str1>", "<str2>"],
  "recommendation": "<approve/return/needs_major_revision>"
}
Return ONLY the JSON, no markdown, no code fences."""
        response = await llm_chat(system, f"Subject: {lesson.get('subject_name','')}\nYear: {lesson.get('year_group','')}\nObjectives: {lesson.get('objectives','')}\nTeacher: {lesson.get('teacher_name','')}\n\nLesson content:\n{content}", provider="gemini"
        )
        import json
        text = response.strip()
        text = re.sub(r'^```(?:json)?\s*\n?', '', text)
        text = re.sub(r'\n?```\s*$', '', text)
        evaluation = json.loads(text)
        now = datetime.now(timezone.utc).isoformat()
        await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
            "evaluation": evaluation, "evaluated_at": now, "evaluated_by": user["user_id"],
        }})
        await log_lesson_action(lesson_id, user["user_id"], user["name"], "evaluated", f"Score: {evaluation.get('overall_score', '?')}/100")
        return evaluation
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Evaluation failed: {str(e)}")


# ─── AI PROCESS UPLOADED FILE ───
@router.post("/{lesson_id}/ai-process-file")
async def ai_process_file(lesson_id: str, data: dict, user: dict = Depends(get_current_user)):
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    file_id = data.get("file_id") or lesson.get("original_file_id")
    if not file_id:
        raise HTTPException(status_code=400, detail="No file to process")
    from storage import get_object
    import tempfile
    tmp_path = None
    try:
        file_meta = await db.file_assets.find_one({"file_id": file_id}, {"_id": 0})
        if not file_meta:
            raise HTTPException(status_code=404, detail="File record not found")
        storage_path = file_meta.get("storage_path", "")
        if not storage_path:
            raise HTTPException(status_code=400, detail="No storage path")
        file_bytes, content_type = get_object(storage_path)
        mime = file_meta.get("content_type", content_type or "application/pdf")
        ext = storage_path.rsplit('.', 1)[-1].lower() if '.' in storage_path else 'pdf'
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=f'.{ext}')
        tmp.write(file_bytes)
        tmp.close()
        tmp_path = tmp.name
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"File read failed: {str(e)}")
    base_prompt = f"""Extract and digitalize this uploaded lesson plan document into a clean, professional HTML format.
Structure it with sections: Learning Objectives, Starter Activity, Main Teaching, Student Activities, Plenary, Assessment, Homework, Resources.
Subject: {lesson.get('subject_name', 'Unknown')}
Year Group: {lesson.get('year_group', '')}
Term: {lesson.get('term', '')}
Return ONLY inner HTML body content. No <!DOCTYPE>, <html>, <head>, <style>, or <body> tags."""
    try:
        from services.ai_service import llm_chat, llm_chat_with_file
        system_for_file = "You are an expert lesson planning assistant for The RightWay Centre of Learning."
        is_pdf = ext == 'pdf' or 'pdf' in mime
        is_image = mime.startswith('image/')
        if is_pdf or is_image:
            response = await llm_chat_with_file(system_for_file, base_prompt, file_path=tmp_path, mime_type=mime)
        else:
            extracted_text = ""
            if ext in ('docx', 'doc'):
                try:
                    from docx import Document
                    doc = Document(tmp_path)
                    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                    for table in doc.tables:
                        for row in table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells)
                            if row_text.strip():
                                paragraphs.append(row_text)
                    extracted_text = "\n".join(paragraphs)
                except Exception:
                    extracted_text = file_bytes.decode('utf-8', errors='ignore')
            else:
                extracted_text = file_bytes.decode('utf-8', errors='ignore')
            if not extracted_text.strip():
                extracted_text = "(Document appears empty)"
            response = await llm_chat(system_for_file, f"{base_prompt}\n\nDocument content:\n\n{extracted_text}", provider="gemini")
        ai_content = _clean_ai_html(response)
        now = datetime.now(timezone.utc).isoformat()
        await db.lessons_v2.update_one({"lesson_id": lesson_id}, {"$set": {
            "ai_version": ai_content, "ai_processed": True, "rich_content": ai_content, "updated_at": now,
        }})
        await log_lesson_action(lesson_id, user["user_id"], user["name"], "ai_digitalized", "File processed by AI")
        return {"ai_content": ai_content}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"AI processing failed: {str(e)}")
    finally:
        try:
            if tmp_path:
                os.unlink(tmp_path)
        except Exception:
            pass


# ─── PDF EXPORT ───
@router.get("/{lesson_id}/export-pdf")
async def export_lesson_pdf(lesson_id: str, request: Request, token: str = Query(None), download: int = Query(0)):
    await _get_user_from_token_or_header(request, token)
    lesson = await db.lessons_v2.find_one({"lesson_id": lesson_id}, {"_id": 0})
    if not lesson:
        raise HTTPException(status_code=404, detail="Not found")
    content = lesson.get("rich_content", "") or lesson.get("content_text", "")
    if not content:
        content = "<p style='color:#888; text-align:center; padding:40px;'>This lesson plan has no content yet. Add content in the editor to generate a full PDF.</p>"
    logo_path = os.path.join(os.path.dirname(__file__), "..", "main_logo.png")
    logo_b64 = ""
    if os.path.exists(logo_path):
        with open(logo_path, "rb") as f:
            logo_b64 = base64.b64encode(f.read()).decode()
    html_doc = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
@page {{ size: A4; margin: 18mm 15mm; }}
body {{ font-family: Arial, Helvetica, sans-serif; font-size: 12px; color: #1a1a1a; line-height: 1.5; }}
.header {{ text-align: center; margin-bottom: 16px; padding-bottom: 10px; border-bottom: 3px solid #1B2A4A; }}
.header img {{ width: 180px; height: auto; margin-bottom: 6px; }}
.header h1 {{ font-size: 15px; color: #1B2A4A; margin: 4px 0; text-transform: uppercase; }}
.header h2 {{ font-size: 13px; color: #2E6CC4; margin: 2px 0; }}
.meta {{ margin-bottom: 14px; font-size: 10px; color: #555; }}
.meta span {{ margin-right: 18px; }}
.meta strong {{ color: #1B2A4A; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 8px; font-size: 11px; }}
thead th {{ background: #1B2A4A; color: white; padding: 7px 8px; text-align: left; font-size: 10px; text-transform: uppercase; border: 1px solid #1B2A4A; }}
tbody td {{ padding: 6px 8px; border: 1px solid #d1d5db; vertical-align: top; }}
h2, h3 {{ color: #1B2A4A; }}
.footer {{ text-align: center; font-size: 8px; color: #999; margin-top: 14px; padding-top: 6px; border-top: 1px solid #e5e7eb; }}
</style></head><body>
<div class="header">
  {"<img src='data:image/png;base64," + logo_b64 + "' />" if logo_b64 else ""}
  <h1>The RightWay Centre of Learning</h1>
  <h2>{lesson.get('title', 'Lesson Plan')}</h2>
</div>
<div class="meta">
  <span>Subject: <strong>{lesson.get('subject_name', '')}</strong></span>
  <span>Year: <strong>{lesson.get('year_group', '')}</strong></span>
  <span>Term: <strong>{lesson.get('term', '')}</strong></span>
  <span>Teacher: <strong>{lesson.get('teacher_name', '')}</strong></span>
  <span>Duration: <strong>{lesson.get('duration_minutes', 45)} min</strong></span>
</div>
{f"<div style='background:#f0f4ff;border:1px solid #c7d2fe;padding:10px;border-radius:6px;margin-bottom:12px;'><strong>Learning Objectives:</strong><br/>{lesson.get('objectives','')}</div>" if lesson.get('objectives') else ""}
{content}
<div class="footer">TCIS Lesson Plan — Generated {datetime.now().strftime('%d %B %Y')} — Confidential</div>
</body></html>"""
    try:
        from weasyprint import HTML
        pdf_bytes = HTML(string=html_doc).write_pdf()
        filename = f"Lesson_{lesson.get('subject_name','').replace(' ','_')}_{lesson.get('year_group','').replace(' ','_')}.pdf"
        disposition = "attachment" if download else "inline"
        return Response(content=pdf_bytes, media_type="application/pdf",
            headers={"Content-Disposition": f"{disposition}; filename={filename}"})
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF generation failed: {str(e)}")


def _clean_ai_html(raw):
    text = raw.strip()
    text = re.sub(r'^```(?:html)?\s*\n?', '', text)
    text = re.sub(r'\n?```\s*$', '', text)
    body_match = re.search(r'<body[^>]*>(.*)</body>', text, re.DOTALL | re.IGNORECASE)
    if body_match:
        text = body_match.group(1).strip()
    else:
        text = re.sub(r'<!DOCTYPE[^>]*>', '', text, flags=re.IGNORECASE)
        text = re.sub(r'</?html[^>]*>', '', text, flags=re.IGNORECASE)
        text = re.sub(r'<head>.*?</head>', '', text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r'<style>.*?</style>', '', text, flags=re.DOTALL | re.IGNORECASE)
    return text.strip()
